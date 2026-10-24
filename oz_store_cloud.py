"""The OZ Store - Cloud Accountant Package Manager (single file)."""
from __future__ import annotations

import io
import re
import urllib.parse
import zipfile
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Side, Border
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------
# Constants shared with the app
# --------------------------------------------------------------------------
ACCOUNTANT_ORDER = [
    "Beverages", "Dried Flower", "Edibles", "Extracts Ingested", "Extracts Inhaled",
    "Infused Pre-Rolls", "Milled Flower", "Pre-Rolls", "Seeds", "Topicals",
    "Vapes - Disposables", "Vapes - 510 Threads", "Vapes - Pods", "Accessories",
    "Apparel / Swag", "Gift Cards", "Other Income", "Raffles", "System Fees", "Uniforms",
]

CHECKLIST_ITEMS = [
    ("brinks", "Brinks Book"),
    ("green1", "Greenline: Inventory Summary"),
    ("green2", "Greenline: Financials"),
    ("green3", "Greenline: Payment Types"),
    ("mon1", "Moneris: Counter Sales"),
    ("mon2", "Moneris: Online Sales"),
    ("ocs1", "OCS: Account Statement"),
    ("ocs2", "OCS: Invoices"),
    ("gas1", "Gas Cards: Details"),
    ("gas2", "Gas Cards: Summary"),
]

EXPENSE_FIELDS = [
    ("rent", "Store Rent & Lease"), ("salaries", "Employee Salaries"),
    ("utilities", "Utilities"), ("advertising", "Advertising"), ("marketing", "Marketing"),
    ("datadeals", "Data Deals & Analytics"), ("insurance", "Insurance"),
    ("garbage", "Garbage Disposal"), ("security", "Security & POS"),
    ("licenses", "Licenses & Compliance"),
]

# Columns that must never be turned into numbers (IDs, SKUs, names...)
_TEXT_COLS = re.compile(
    r"\b(id|sku|upc|barcode|name|product|description|category|brand|supplier|vendor|store|date|code|plu)\b",
    re.I,
)
MAX_SHEET_ROWS = 50_000

# --------------------------------------------------------------------------
# Small helpers
# --------------------------------------------------------------------------
def parse_money(value) -> float:
    """'$1,234.50', '(200)', '1 234,50' -> float. Blank/invalid -> 0.0."""
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return 0.0 if value != value else float(value)
    s = str(value).strip()
    if not s:
        return 0.0
    neg = (s.startswith("(") and s.endswith(")")) or s.startswith("-")
    s = s.replace(" ", "").replace("\u00a0", "")
    if "," in s and "." not in s and re.search(r",\d{1,2}$", s):
        s = s.replace(",", ".")  # European / Quebec decimal comma
    s = re.sub(r"[^\d.]", "", s)
    try:
        x = float(s)
    except ValueError:
        return 0.0
    return -x if neg else x


def total_expenses(payload: dict) -> float:
    return sum(parse_money(v) for v in (payload.get("expenses") or {}).values())


def missing_docs(payload: dict) -> list[str]:
    chk = payload.get("checklist") or {}
    files = payload.get("files") or {}
    return [label for key, label in CHECKLIST_ITEMS if not (chk.get(key) or files.get(key))]


def checklist_done(payload: dict) -> int:
    return len(CHECKLIST_ITEMS) - len(missing_docs(payload))


# --------------------------------------------------------------------------
# Reading uploaded files
# --------------------------------------------------------------------------
def _read_ragged(text: str) -> pd.DataFrame:
    """Fallback for exports with title lines above the header (rows of different widths)."""
    import csv

    rows = list(csv.reader(io.StringIO(text)))
    width = max((len(r) for r in rows), default=0)
    if width == 0:
        return pd.DataFrame()
    rows = [r + [""] * (width - len(r)) for r in rows]
    fill = [sum(1 for v in r if str(v).strip()) for r in rows]
    best = max(fill[:25])
    hdr = next(i for i, n in enumerate(fill[:25]) if n == best)
    df = pd.DataFrame(rows[hdr + 1:], columns=[str(v).strip() or f"Column {j + 1}" for j, v in enumerate(rows[hdr])])
    return df.replace(r"^\s*$", pd.NA, regex=True)


def _read_csv(data: bytes) -> pd.DataFrame:
    last_err = None
    for enc in ("utf-8-sig", "cp1252"):
        try:
            text = data.decode(enc)
        except UnicodeDecodeError as e:
            last_err = e
            continue
        try:
            df = pd.read_csv(io.StringIO(text))
            if df.shape[1] == 1:  # maybe ; or tab separated
                alt = pd.read_csv(io.StringIO(text), sep=None, engine="python")
                if alt.shape[1] > 1:
                    df = alt
            return df
        except (pd.errors.ParserError, pd.errors.EmptyDataError):
            return _read_ragged(text)
    raise ValueError(f"Could not decode CSV: {last_err}")


def read_table(name: str, data: bytes) -> pd.DataFrame | None:
    """CSV / Excel -> DataFrame. Anything else (PDF, images) -> None."""
    ext = name.lower().rsplit(".", 1)[-1] if "." in name else ""
    if ext in ("csv", "txt"):
        df = _read_csv(data)
    elif ext in ("xlsx", "xlsm", "xls"):
        df = pd.read_excel(io.BytesIO(data))
    else:
        return None
    return df.dropna(how="all").dropna(axis=1, how="all").reset_index(drop=True)


def _promote_header(df: pd.DataFrame, needle: str = "category", scan: int = 25) -> pd.DataFrame:
    """Some exports have title lines above the real header; find it."""
    if any(str(c).strip().lower() == needle for c in df.columns):
        return df
    for i in range(min(scan, len(df))):
        row = [str(v).strip() for v in df.iloc[i].tolist()]
        if needle in [r.lower() for r in row]:
            out = df.iloc[i + 1:].copy()
            out.columns = row
            return out.reset_index(drop=True)
    return df


def _coerce_numeric(df: pd.DataFrame) -> pd.DataFrame:
    for c in df.columns:
        if pd.api.types.is_numeric_dtype(df[c]) or _TEXT_COLS.search(str(c)):
            continue
        s = df[c].astype(str).str.strip()
        filled = df[c].notna() & (s != "") & (s.str.lower() != "nan")
        if not filled.any():
            continue
        cleaned = (
            s.str.replace(r"^\((.*)\)$", r"-\1", regex=True)
            .str.replace(r"[$,%\s]", "", regex=True)
        )
        num = pd.to_numeric(cleaned, errors="coerce")
        if num[filled].notna().mean() >= 0.9:
            df[c] = num
    return df


def numeric_cols(df: pd.DataFrame) -> list[str]:
    return [
        c for c in df.columns
        if pd.api.types.is_numeric_dtype(df[c])
        and not pd.api.types.is_bool_dtype(df[c])
        and not _TEXT_COLS.search(str(c))
    ]


def clean_inventory(df: pd.DataFrame) -> tuple[pd.DataFrame, list[str]]:
    """Order categories like the accountant wants, drop subtotal rows, fix numbers."""
    notes: list[str] = []
    df = df.dropna(how="all").dropna(axis=1, how="all").copy()
    df = _promote_header(df)
    df.columns = [str(c).strip() for c in df.columns]

    cat_col = next((c for c in df.columns if c.lower() == "category"), None)
    if cat_col is None:
        notes.append("No 'Category' column found, so rows were kept in their original order.")
        return _coerce_numeric(df), notes
    if cat_col != "Category":
        df = df.rename(columns={cat_col: "Category"})

    cats = df["Category"].astype(str).str.strip()
    drop = cats.str.startswith(">") | cats.str.lower().isin(
        ["", "nan", "none", "total", "totals", "grand total"]
    )
    df, cats = df[~drop].copy(), cats[~drop]

    rank = {c.casefold(): i for i, c in enumerate(ACCOUNTANT_ORDER)}
    order = cats.str.casefold().map(rank)
    unknown = sorted(cats[order.isna()].unique())
    if unknown:
        notes.append("Categories not in the accountant order (placed at the end): " + ", ".join(unknown))
    df["Category"] = cats
    df["_o"] = order.fillna(len(ACCOUNTANT_ORDER))
    df = df.sort_values("_o", kind="stable").drop(columns="_o").reset_index(drop=True)
    return _coerce_numeric(df), notes


# --------------------------------------------------------------------------
# Excel styling helpers
# --------------------------------------------------------------------------
_FONT = "Arial"
_HDR_FILL = PatternFill("solid", fgColor="065F46")
_HDR_FONT = Font(name=_FONT, bold=True, color="FFFFFF", size=10)
_TOT_FILL = PatternFill("solid", fgColor="D1FAE5")
_BAD_FILL = PatternFill("solid", fgColor="FEE2E2")
_BODY = Font(name=_FONT, size=10)
_BOLD = Font(name=_FONT, size=10, bold=True)
_THIN = Border(bottom=Side(style="thin", color="E5E7EB"))
MONEY = '#,##0.00;[Red]-#,##0.00'
INT = '#,##0;[Red]-#,##0'


def _cell_value(v):
    if v is None:
        return None
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    if hasattr(v, "item"):
        return v.item()
    return v


def _title(ws, row: int, text: str, size: int = 12):
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(name=_FONT, bold=True, size=size, color="065F46")


def _header(ws, row: int, headers: list[str]):
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=j, value=h)
        c.font, c.fill = _HDR_FONT, _HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _put(ws, row: int, col: int, value, *, bold=False, fill=None):
    c = ws.cell(row=row, column=col, value=_cell_value(value))
    c.font = _BOLD if bold else _BODY
    c.border = _THIN
    if fill:
        c.fill = fill
    if isinstance(c.value, bool):
        pass
    elif isinstance(c.value, int):
        c.number_format = INT
    elif isinstance(c.value, float):
        c.number_format = MONEY
    return c


def _write_df(ws, df: pd.DataFrame, row: int = 1, totals: bool = True) -> int:
    """Write a styled table; returns the next free row."""
    df = df.head(MAX_SHEET_ROWS)
    _header(ws, row, [str(c) for c in df.columns])
    for i, rec in enumerate(df.itertuples(index=False), start=row + 1):
        for j, v in enumerate(rec, 1):
            _put(ws, i, j, v)
    r = row + len(df) + 1
    nums = numeric_cols(df) if totals else []
    if nums:
        for j, c in enumerate(df.columns, 1):
            if j == 1:
                _put(ws, r, j, "TOTAL", bold=True, fill=_TOT_FILL)
            elif c in nums:
                _put(ws, r, j, float(df[c].sum(skipna=True)), bold=True, fill=_TOT_FILL)
            else:
                _put(ws, r, j, None, fill=_TOT_FILL)
        r += 1
    ws.freeze_panes = ws.cell(row=row + 1, column=1)
    return r + 1


def _block(ws, row: int, title: str, headers: list[str], rows: list[list], bad_when=None) -> int:
    _title(ws, row, title)
    _header(ws, row + 1, headers)
    for i, rec in enumerate(rows, start=row + 2):
        bad = bad_when(rec) if bad_when else False
        for j, v in enumerate(rec, 1):
            _put(ws, i, j, v, fill=_BAD_FILL if bad else None)
    return row + len(rows) + 3


def _autowidth(ws, cap: int = 48):
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is None:
                continue
            n = len(f"{c.value:,.2f}") if isinstance(c.value, float) else len(str(c.value))
            widths[c.column] = max(widths.get(c.column, 0), min(n, cap))
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(10, w + 2)


def _sheet_name(name: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]\*\?/\\:]", "", name).strip()[:31] or "Sheet"
    out, n = base, 2
    while out.lower() in used:
        suffix = f" ({n})"
        out, n = base[: 31 - len(suffix)] + suffix, n + 1
    used.add(out.lower())
    return out


def _to_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _inventory_frame(payload: dict) -> pd.DataFrame | None:
    recs = payload.get("inventory_data")
    if not recs:
        return None
    df = pd.DataFrame(recs)
    df = df.replace({"": None})
    return _coerce_numeric(df)


def _num(v) -> float:
    try:
        return float(re.sub(r"[^0-9.\-]", "", str(v))) if v is not None else 0.0
    except ValueError:
        return 0.0


def financials(payload: dict) -> dict:
    """Same maths as the on-screen Financial Summary (sales = 'sold' column, COGS = 'cost' column, 80% fallback)."""
    rows = payload.get("inventory_data") or []
    sales = cogs = 0.0
    if rows:
        headers = list(rows[0].keys())
        sold_h = next((h for h in headers if h and "sold" in str(h).lower()), None)
        cost_h = next((h for h in headers if h and "cost" in str(h).lower()), None)
        for r in rows:
            if sold_h and r.get(sold_h) is not None:
                sales += _num(r.get(sold_h))
            else:
                keys = [k for k in r if k != "Category"]
                if keys:
                    sales += _num(r.get(keys[0]))
            if cost_h and r.get(cost_h) is not None:
                cogs += _num(r.get(cost_h))
    estimated = False
    if cogs == 0 and sales > 0:
        cogs, estimated = sales * 0.80, True
    gross = sales - cogs
    exp = total_expenses(payload)
    noi = gross - exp
    return {
        "sales": sales, "cogs": cogs, "gross": gross, "expenses": exp, "noi": noi,
        "gross_margin": (gross / sales * 100) if sales > 0 else 0.0,
        "op_margin": (noi / sales * 100) if sales > 0 else 0.0,
        "cogs_estimated": estimated,
    }


# --------------------------------------------------------------------------
# Report builders
# --------------------------------------------------------------------------
def inventory_workbook(df: pd.DataFrame, title: str = "Inventory") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = _sheet_name(title, set())
    _write_df(ws, df)
    _autowidth(ws)
    return _to_bytes(wb)


def build_store_workbook(
    store: str, period: str, payload: dict, tables: dict[str, tuple[str, pd.DataFrame]],
    doc_log: list[list] | None = None,
) -> bytes:
    """One store, one month: Summary + Inventory + one sheet per tabular document + Documents log."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    used = {"summary"}

    ws["A1"] = f"The OZ Store - {store}"
    ws["A1"].font = Font(name=_FONT, bold=True, size=16, color="065F46")
    ws["A2"] = f"Period: {period}"
    ws["A3"] = f"Generated: {datetime.now():%Y-%m-%d %H:%M}"
    for a_ in ("A2", "A3"):
        ws[a_].font = _BODY
    missing = missing_docs(payload)
    ws["A4"] = ("Missing documents: " + ", ".join(missing)) if missing else "All documents received"
    ws["A4"].font = Font(name=_FONT, bold=True, size=10, color="B91C1C" if missing else "047857")

    files = payload.get("files") or {}
    notes = payload.get("notes") or {}
    chk = payload.get("checklist") or {}

    def names(key):
        v = files.get(key) or []
        if isinstance(v, dict):
            v = [v]
        return ", ".join(f.get("name", "") for f in v)

    rows = [
        [label, "Received" if (chk.get(key) or files.get(key)) else "MISSING", names(key), notes.get(key, "")]
        for key, label in CHECKLIST_ITEMS
    ]
    r = _block(ws, 6, "Document checklist", ["Document", "Status", "File(s)", "Note"], rows,
               bad_when=lambda rec: rec[1] == "MISSING")

    cooh_text = payload.get("cooh_text", "")
    cooh_amt = parse_money(cooh_text) if str(cooh_text).strip() else None
    r = _block(ws, r, "Cash & chargebacks", ["Item", "Value"], [
        ["Cash on hand", cooh_amt if cooh_amt is not None else "Not entered"],
        ["Cash on hand date", payload.get("cooh_date") or "Not entered"],
        ["Moneris chargebacks", payload.get("chargebacks_text") or "Nothing to report."],
    ])

    exp = payload.get("expenses") or {}
    exp_rows = [[label, parse_money(exp.get(key))] for key, label in EXPENSE_FIELDS]
    exp_rows.append(["TOTAL", total_expenses(payload)])
    r = _block(ws, r, "Operating expenses", ["Expense", "Amount"], exp_rows)

    fin = financials(payload)
    if fin["sales"] > 0:
        r = _block(ws, r, "Financial summary", ["Measure", "Value"], [
            ["Total sales (inventory 'sold' column)", fin["sales"]],
            ["Cost of goods sold" + (" (estimated at 80%)" if fin["cogs_estimated"] else ""), fin["cogs"]],
            ["Gross profit", fin["gross"]],
            ["Gross margin %", round(fin["gross_margin"], 1)],
            ["Operating expenses", fin["expenses"]],
            ["Net operating income", fin["noi"]],
            ["Operating margin %", round(fin["op_margin"], 1)],
        ])

    inv = _inventory_frame(payload)
    if inv is not None:
        nums = numeric_cols(inv)[:12]
        inv_rows = [["Categories", int(inv["Category"].nunique()) if "Category" in inv else len(inv)],
                    ["Line items", len(inv)]]
        inv_rows += [[c if c.lower().startswith("total") else f"Total {c}", float(inv[c].sum(skipna=True))] for c in nums]
        r = _block(ws, r, "Inventory totals", ["Measure", "Value"], inv_rows)
        ws2 = wb.create_sheet(_sheet_name("Inventory", used))
        _write_df(ws2, inv)
        _autowidth(ws2)

    for key, (label, df) in tables.items():
        if len(df) > MAX_SHEET_ROWS:
            df = df.head(MAX_SHEET_ROWS)
        sheet = wb.create_sheet(_sheet_name(label, used))
        _write_df(sheet, df, totals=True)
        _autowidth(sheet)

    if doc_log:
        wd = wb.create_sheet(_sheet_name("Documents", used))
        _header(wd, 1, ["Category", "File", "Size (KB)", "In this report"])
        for i, rec in enumerate(doc_log, start=2):
            for j, v in enumerate(rec, 1):
                _put(wd, i, j, v, fill=_BAD_FILL if "NOT FOUND" in str(rec[-1]) else None)
        _autowidth(wd)

    _autowidth(ws)
    return _to_bytes(wb)


def build_consolidated_workbook(period: str, payloads: dict[str, dict]) -> bytes:
    """All stores for one month: Overview, Expenses matrix, All Inventory, Category Totals."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Overview"
    used = {"overview"}
    ws["A1"] = f"The OZ Store - All stores - {period}"
    ws["A1"].font = Font(name=_FONT, bold=True, size=16, color="065F46")
    ws["A2"] = f"Generated: {datetime.now():%Y-%m-%d %H:%M}"
    ws["A2"].font = _BODY

    rows = []
    for store, p in payloads.items():
        cooh = str(p.get("cooh_text", "")).strip()
        fin = financials(p)
        rows.append([
            store,
            f"{checklist_done(p)}/{len(CHECKLIST_ITEMS)}",
            parse_money(cooh) if cooh else "Not entered",
            p.get("cooh_date") or "",
            p.get("chargebacks_text") or "",
            total_expenses(p),
            fin["sales"],
            fin["noi"],
            ", ".join(missing_docs(p)) or "-",
        ])
    if rows:
        money = [r_ for r_ in rows]
        rows.append(["ALL STORES", "", "", "", "",
                     sum(r_[5] for r_ in money), sum(r_[6] for r_ in money), sum(r_[7] for r_ in money), ""])
    _block(ws, 4, "Store overview",
           ["Store", "Checklist", "Cash on hand", "Cash date", "Chargebacks", "Total expenses", "Total sales",
            "Net operating income", "Missing"],
           rows, bad_when=lambda rec: rec[8] not in ("-", ""))
    _autowidth(ws)

    stores = list(payloads)
    wx = wb.create_sheet(_sheet_name("Expenses", used))
    _header(wx, 1, ["Expense"] + stores + ["All stores"])
    for i, (key, label) in enumerate(EXPENSE_FIELDS, start=2):
        vals = [parse_money((payloads[s].get("expenses") or {}).get(key)) for s in stores]
        _put(wx, i, 1, label)
        for j, v in enumerate(vals, 2):
            _put(wx, i, j, v)
        _put(wx, i, len(stores) + 2, sum(vals), bold=True)
    tr = len(EXPENSE_FIELDS) + 2
    _put(wx, tr, 1, "TOTAL", bold=True, fill=_TOT_FILL)
    for j, s in enumerate(stores, 2):
        _put(wx, tr, j, total_expenses(payloads[s]), bold=True, fill=_TOT_FILL)
    _put(wx, tr, len(stores) + 2, sum(total_expenses(payloads[s]) for s in stores), bold=True, fill=_TOT_FILL)
    wx.freeze_panes = "B2"
    _autowidth(wx)

    frames = []
    for s in stores:
        f = _inventory_frame(payloads[s])
        if f is not None:
            f.insert(0, "Store", s)
            frames.append(f)
    if frames:
        allinv = pd.concat(frames, ignore_index=True)
        wa = wb.create_sheet(_sheet_name("All Inventory", used))
        _write_df(wa, allinv, totals=True)
        _autowidth(wa)
        nums = numeric_cols(allinv)
        if "Category" in allinv and nums:
            tot = allinv.groupby("Category", sort=False)[nums].sum()
            rank = {c.casefold(): i for i, c in enumerate(ACCOUNTANT_ORDER)}
            tot = tot.assign(_o=[rank.get(str(c).casefold(), 999) for c in tot.index]) \
                     .sort_values("_o", kind="stable").drop(columns="_o").reset_index()
            wc = wb.create_sheet(_sheet_name("Category Totals", used))
            _write_df(wc, tot, totals=True)
            _autowidth(wc)
    return _to_bytes(wb)


def build_zip(entries: list[tuple[str, bytes]]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for arcname, data in entries:
            z.writestr(arcname, data)
    return buf.getvalue()


# ==========================================================================


# ==========================================================================
# THE PAGE (your original Tailwind / emerald interface, unchanged in look)
# ==========================================================================
HTML_CODE = r"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>The OZ Store - Executive Financial Report & Package Manager</title>
    <script src="https://cdn.tailwindcss.com"></script>
    <script src="https://cdnjs.cloudflare.com/ajax/libs/PapaParse/5.4.1/papaparse.min.js"></script>
    <script src="https://cdnjs.cloudflare.com/ajax/libs/xlsx/0.18.5/xlsx.full.min.js"></script>
    <script src="https://cdnjs.cloudflare.com/ajax/libs/jszip/3.10.1/jszip.min.js"></script>
    <script src="https://cdnjs.cloudflare.com/ajax/libs/FileSaver.js/2.0.5/FileSaver.min.js"></script>
</head>
<body class="bg-zinc-100 text-zinc-900 min-h-screen flex flex-col items-center py-8 px-4 font-sans">

    <div class="max-w-7xl w-full bg-white rounded-2xl shadow-xl overflow-hidden border border-zinc-200">

        <header class="bg-emerald-800 text-white p-6 md:p-8 flex flex-col md:flex-row justify-between items-center gap-4 border-b border-emerald-900">
            <div class="flex items-center gap-4">
                <div class="bg-white p-2.5 rounded-xl shadow-md flex items-center justify-center">
                    <span class="text-2xl">🌿</span>
                </div>
                <div>
                    <h1 class="text-xl md:text-2xl font-extrabold text-white tracking-tight">The OZ Store</h1>
                    <p class="text-emerald-100/80 text-xs mt-0.5">Executive Financial & Accountant Document Package Manager</p>
                </div>
            </div>

            <div class="flex items-center gap-4">
                <div class="flex items-center gap-3 bg-white text-zinc-900 px-4 py-2 rounded-xl shadow-md border border-emerald-700">
                    <span class="text-xs font-bold text-emerald-800 uppercase tracking-wider">Period:</span>
                    <select id="select-month" onchange="updatePeriod()" class="bg-zinc-50 text-zinc-900 text-xs font-semibold px-2.5 py-1.5 rounded-lg border border-zinc-300 focus:outline-none focus:border-emerald-600">
                        <option value="January">January</option>
                        <option value="February">February</option>
                        <option value="March">March</option>
                        <option value="April">April</option>
                        <option value="May">May</option>
                        <option value="June">June</option>
                        <option value="July">July</option>
                        <option value="August">August</option>
                        <option value="September">September</option>
                        <option value="October" selected>October</option>
                        <option value="November">November</option>
                        <option value="December">December</option>
                    </select>
                    <select id="select-year" onchange="updatePeriod()" class="bg-zinc-50 text-zinc-900 text-xs font-semibold px-2.5 py-1.5 rounded-lg border border-zinc-300 focus:outline-none focus:border-emerald-600">
                        <option value="2025">2025</option>
                        <option value="2026" selected>2026</option>
                        <option value="2027">2027</option>
                        <option value="2028">2028</option>
                    </select>
                </div>
                <button onclick="navigateParent('?auth_user=' + encodeURIComponent(currentUsername) + '&action=sign_out')" class="bg-emerald-900 hover:bg-emerald-950 text-emerald-100 px-3.5 py-2 rounded-xl text-xs font-semibold border border-emerald-700 transition">
                    Sign Out (REPLACE_CURRENT_NAME)
                </button>
            </div>
        </header>

        <div class="p-6 md:p-8">

            <div class="flex flex-col md:flex-row justify-between items-center mb-6 gap-4 pb-4 border-b border-zinc-200">
                <div class="flex items-center gap-2 overflow-x-auto pb-2 scrollbar-thin w-full md:w-auto" id="store-tabs"></div>
                <div class="flex items-center gap-3">
                    REPLACE_ADD_USER_BTN_HTML
                    <div class="text-xs text-zinc-500 font-medium bg-zinc-50 px-3 py-1.5 rounded-lg border border-zinc-200">
                        👤 Logged in as: <span class="font-bold text-emerald-800">REPLACE_CURRENT_NAME</span>
                    </div>
                </div>
            </div>

            <!-- COLLAPSIBLE CREDENTIAL MANAGER (ADMIN ONLY) -->
            <details id="credential-manager-view" class="REPLACE_ADMIN_HIDDEN_CLASS mb-8 bg-zinc-50 rounded-2xl border border-zinc-300 shadow-xs overflow-hidden">
                <summary class="p-4 bg-emerald-50 hover:bg-emerald-100 cursor-pointer font-bold text-emerald-900 text-xs flex items-center justify-between transition">
                    <span class="flex items-center gap-2"><span>🔐</span> Employee & Store Login Credentials Directory (Click to Edit)</span>
                    <span class="text-[11px] text-emerald-700 underline font-normal">Expand / Collapse</span>
                </summary>
                <div class="p-5 border-t border-zinc-200 bg-white space-y-4">
                    <div class="bg-white rounded-xl border border-zinc-200 overflow-hidden shadow-xs">
                        <table class="w-full text-left border-collapse text-xs">
                            <thead class="bg-emerald-50 text-emerald-900 border-b border-emerald-200">
                                <tr>
                                    <th class="p-3 font-semibold">Username</th>
                                    <th class="p-3 font-semibold">Display Name</th>
                                    <th class="p-3 font-semibold">Role</th>
                                    <th class="p-3 font-semibold">Password</th>
                                    <th class="p-3 font-semibold text-right">Actions</th>
                                </tr>
                            </thead>
                            <tbody class="divide-y divide-zinc-100 text-zinc-700">
                                REPLACE_CRED_TABLE_ROWS
                            </tbody>
                        </table>
                    </div>

                    <div id="user-edit-form-box" class="hidden bg-emerald-50/60 p-4 rounded-xl border border-emerald-300 mt-4">
                        <h4 class="font-bold text-emerald-900 text-xs mb-3" id="edit-form-title">Edit User Credential</h4>
                        <form onsubmit="submitUserEdit(event)" class="grid grid-cols-1 md:grid-cols-5 gap-3 text-xs">
                            <input type="hidden" id="edit-orig-username">
                            <div>
                                <label class="block font-semibold mb-1 text-zinc-700">Username</label>
                                <input type="text" id="edit-username" required class="w-full border border-zinc-300 rounded px-2 py-1.5 bg-white">
                            </div>
                            <div>
                                <label class="block font-semibold mb-1 text-zinc-700">Display Name</label>
                                <input type="text" id="edit-name" required class="w-full border border-zinc-300 rounded px-2 py-1.5 bg-white">
                            </div>
                            <div>
                                <label class="block font-semibold mb-1 text-zinc-700">Role</label>
                                <select id="edit-role" class="w-full border border-zinc-300 rounded px-2 py-1.5 bg-white">
                                    <option value="store">Store</option>
                                    <option value="accountant">Accountant</option>
                                    <option value="admin">Admin</option>
                                </select>
                            </div>
                            <div>
                                <label class="block font-semibold mb-1 text-zinc-700">Password</label>
                                <input type="text" id="edit-password" required class="w-full border border-zinc-300 rounded px-2 py-1.5 bg-white">
                            </div>
                            <div class="flex items-end gap-2">
                                <button type="submit" class="bg-emerald-700 hover:bg-emerald-800 text-white font-bold px-3 py-1.5 rounded transition w-full">Save</button>
                                <button type="button" onclick="document.getElementById('user-edit-form-box').classList.add('hidden')" class="bg-zinc-200 text-zinc-700 px-3 py-1.5 rounded transition">Cancel</button>
                            </div>
                        </form>
                    </div>
                </div>
            </details>

            <div class="flex justify-between items-center mb-6 bg-emerald-50/60 p-4 rounded-xl border border-emerald-200 shadow-xs">
                <div>
                    <span class="text-xs uppercase tracking-wider text-emerald-800 font-bold block mb-0.5">Active Workspace</span>
                    <h2 class="text-xl font-extrabold text-zinc-900" id="workspace-store-name">Oz 1</h2>
                </div>
                <div class="flex items-center gap-3 text-xs text-zinc-600">
                    <button type="button" onclick="saveAllData(false)" class="bg-emerald-700 hover:bg-emerald-800 text-white px-4 py-1.5 rounded-full font-semibold shadow-xs border border-emerald-600 transition cursor-pointer">
                        💾 Save All to Cloud
                    </button>
                    <span id="active-period-label" class="bg-white text-emerald-800 border border-emerald-300 px-3 py-1 rounded-full font-semibold shadow-xs">October 2026</span>
                    <span id="store-completion-badge" class="bg-amber-50 text-amber-800 border border-amber-300 px-3 py-1 rounded-full font-semibold shadow-xs">Checklist: 0/10 Completed</span>
                </div>
            </div>

            <!-- SIDE-BY-SIDE SECTIONS -->
            <div class="grid grid-cols-1 xl:grid-cols-12 gap-6 mb-8 REPLACE_ACCOUNTANT_HIDDEN_CLASS">

                <!-- LEFT SIDE: Inventory Report Formatter -->
                <div class="xl:col-span-5 bg-zinc-50 rounded-xl p-5 border border-zinc-200 flex flex-col justify-between shadow-xs">
                    <div>
                        <div class="flex justify-between items-center mb-3">
                            <h3 class="text-base font-bold text-emerald-800 flex items-center gap-2">Inventory Report Formatter</h3>
                            <span id="inv-status-badge" class="text-[11px] text-zinc-600 bg-white border border-zinc-300 px-2 py-0.5 rounded shadow-xs">Not Loaded</span>
                        </div>
                        <p class="text-zinc-600 text-xs mb-4">Upload raw monthly CSV. Rows are sorted and automatically pushed to Greenline Inventory Summary.</p>

                        <div id="drop-zone" class="border-2 border-dashed border-emerald-400 hover:border-emerald-700 rounded-xl p-5 text-center cursor-pointer transition-all bg-white group mb-4 shadow-xs">
                            <input type="file" id="file-input" accept=".csv" class="hidden">
                            <svg class="w-8 h-8 mx-auto text-emerald-700 mb-2 group-hover:scale-110 transition-transform" fill="none" stroke="currentColor" viewBox="0 0 24 24">
                                <path stroke-linecap="round" stroke-linejoin="round" stroke-width="2" d="M7 16a4 4 0 01-.88-7.903A5 5 0 1115.9 6L16 6a5 5 0 011 9.9M15 13l-3-3m0 0l-3 3m3-3v12"></path>
                            </svg>
                            <p class="text-xs font-medium text-zinc-700">Drag & drop store CSV here</p>
                            <p class="text-[11px] text-emerald-700 underline mt-0.5">or click to browse</p>
                        </div>
                        <div id="csv-file-info" class="hidden mb-4 bg-emerald-50 border border-emerald-200 rounded-xl p-3 flex items-center justify-between">
                            <div class="flex items-center gap-2">
                                <span class="text-emerald-700">📄</span>
                                <span id="csv-file-name" class="text-xs font-medium text-emerald-800"></span>
                            </div>
                            <button type="button" onclick="clearCSVUpload()" class="bg-red-100 hover:bg-red-200 text-red-700 text-xs font-bold px-3 py-1.5 rounded-lg border border-red-300 transition">
                                🗑️ Delete & Re-upload
                            </button>
                        </div>

                        <div class="bg-white rounded-lg border border-zinc-200 overflow-hidden max-h-56 overflow-y-auto mb-4 shadow-xs">
                            <table class="w-full text-left border-collapse text-[11px]">
                                <thead class="bg-emerald-50 text-emerald-900 sticky top-0 z-10 border-b border-emerald-200">
                                    <tr id="table-header"><th class="p-2 font-semibold">Category</th><th class="p-2 font-semibold">Opening</th><th class="p-2 font-semibold">Sold</th><th class="p-2 font-semibold">Closing</th></tr>
                                </thead>
                                <tbody id="table-body" class="divide-y divide-zinc-100 text-zinc-700"><tr><td colspan="4" class="p-4 text-center text-zinc-400">No report uploaded yet.</td></tr></tbody>
                            </table>
                        </div>
                    </div>

                    <div id="action-area" class="hidden">
                        <button id="download-excel-btn" class="w-full bg-emerald-700 hover:bg-emerald-800 text-white font-bold py-2.5 px-4 rounded-xl shadow-md transition-all flex items-center justify-center gap-2 text-xs">
                            Download Formatted Excel (.xlsx)
                        </button>
                    </div>
                </div>

                <!-- RIGHT SIDE: Monthly Accountant Document Checklist -->
                <div class="xl:col-span-7 bg-zinc-50 rounded-xl p-5 border border-zinc-200 flex flex-col justify-between shadow-xs">
                    <div>
                        <div class="flex justify-between items-center mb-3">
                            <h3 class="text-base font-bold text-emerald-800">Monthly Accountant Package Checklist & Folders</h3>
                        </div>
                        <p class="text-zinc-600 text-xs mb-4">Input cash on hand, browse supporting files, and add notes. Click <b>Save All to Cloud</b> when ready.</p>

                        <div class="space-y-3.5 max-h-[480px] overflow-y-auto pr-2" id="checklist-container">

                            <!-- 1. Cash on Hand -->
                            <div class="bg-white p-3 rounded-xl border border-zinc-200 space-y-2 shadow-xs">
                                <div class="flex justify-between items-center">
                                    <label class="block text-xs font-bold text-emerald-800">Cash On Hand - End of Day</label>
                                    <input type="date" id="cooh-date" onchange="updateCompletionBadge(); saveLocalOnly();" class="bg-white text-zinc-700 text-[11px] px-2 py-1 rounded border border-zinc-300">
                                </div>
                                <div class="grid grid-cols-1 md:grid-cols-2 gap-2">
                                    <input type="text" id="cooh-text" placeholder="Enter amount e.g. $24,521.10" oninput="syncNotes('cooh', this.value)" class="w-full bg-white border border-zinc-300 rounded-lg px-3 py-2 text-xs text-zinc-800">
                                    <div class="flex flex-col gap-1">
                                        <input type="file" id="cooh-file" multiple onchange="handleItemFiles('cooh', this)" class="hidden">
                                        <div class="flex gap-1">
                                            <button type="button" onclick="openUploadPanel('cooh')" class="flex-1 bg-emerald-50 hover:bg-emerald-100 text-emerald-800 text-xs font-medium py-2 px-3 rounded-lg border border-emerald-300 transition truncate">📎 Browse & Upload Document</button>
                                            <button type="button" onclick="clearUploadField('cooh')" class="bg-red-50 hover:bg-red-100 text-red-600 text-xs font-bold px-2 py-2 rounded-lg border border-red-200 transition" title="Clear uploaded files">🗑️</button>
                                        </div>
                                    </div>
                                </div>
                                <div id="files-cooh" class="flex flex-wrap gap-1.5 mt-1 text-[11px] text-zinc-600">No files uploaded</div>
                            </div>

                            <!-- 2. Brinks Book -->
                            <div class="bg-white p-3 rounded-xl border border-zinc-200 space-y-2 shadow-xs">
                                <div class="flex items-center justify-between">
                                    <label class="flex items-center gap-2 text-xs font-bold text-emerald-800 cursor-pointer">
                                        <input type="checkbox" id="brinks-chk" onchange="updateCompletionBadge(); saveLocalOnly();" class="accent-emerald-600 w-4 h-4 rounded"> Brinks Book (Folder)
                                    </label>
                                    <input type="date" id="brinks-date" onchange="updateCompletionBadge(); saveLocalOnly();" class="bg-white text-zinc-700 text-[11px] px-2 py-1 rounded border border-zinc-300">
                                </div>
                                <div class="pl-5 space-y-2">
                                    <div class="grid grid-cols-1 md:grid-cols-2 gap-2">
                                        <input type="text" id="brinks-note" placeholder="Add notes..." oninput="syncNotes('brinks', this.value)" class="w-full bg-white border border-zinc-300 rounded-lg px-3 py-2 text-xs text-zinc-800">
                                        <div class="flex flex-col gap-1">
                                            <input type="file" id="brinks-file" multiple onchange="handleItemFiles('brinks', this)" class="hidden">
                                            <div class="flex gap-1">
                                                <button type="button" onclick="openUploadPanel('brinks')" class="flex-1 bg-emerald-50 hover:bg-emerald-100 text-emerald-800 text-xs font-medium py-2 px-3 rounded-lg border border-emerald-300 transition truncate">📎 Browse & Upload Photos</button>
                                                <button type="button" onclick="clearUploadField('brinks')" class="bg-red-50 hover:bg-red-100 text-red-600 text-xs font-bold px-2 py-2 rounded-lg border border-red-200 transition" title="Clear uploaded files">🗑️</button>
                                            </div>
                                        </div>
                                    </div>
                                    <div id="files-brinks" class="flex flex-wrap gap-1.5 mt-1 text-[11px] text-zinc-600">No files uploaded</div>
                                </div>
                            </div>

                            <!-- 3. Greenline Reports -->
                            <div class="bg-white p-3 rounded-xl border border-zinc-200 space-y-2.5 shadow-xs">
                                <span class="text-xs font-bold text-emerald-800 block">Greenline Reports</span>
                                <div class="pl-3 space-y-2.5 bg-zinc-50 p-2.5 rounded-lg border border-zinc-200">
                                    <div class="flex justify-between items-center">
                                        <label class="text-xs font-medium flex items-center gap-1.5"><input type="checkbox" id="green1-chk" onchange="updateCompletionBadge(); saveLocalOnly();"> Inventory Summary (Auto-pushed from Formatter)</label>
                                        <input type="file" id="green1-file" multiple onchange="handleItemFiles('green1', this)" class="hidden">
                                        <div class="flex gap-1">
                                            <button type="button" onclick="openUploadPanel('green1')" class="bg-white hover:bg-zinc-100 text-zinc-700 text-[11px] py-1 px-2.5 rounded border border-zinc-300">📎 Browse</button>
                                            <button type="button" onclick="clearUploadField('green1')" class="bg-red-50 hover:bg-red-100 text-red-600 text-[11px] font-bold px-1.5 py-1 rounded border border-red-200 transition" title="Clear uploaded files">🗑️</button>
                                        </div>
                                    </div>
                                    <div id="files-green1" class="flex flex-wrap gap-1.5 mt-1 text-[11px] text-zinc-600">No files</div>
                                    <input type="text" id="green1-note" placeholder="Notes..." oninput="syncNotes('green1', this.value)" class="w-full bg-white border border-zinc-300 rounded px-2 py-1 text-xs">

                                    <div class="flex justify-between items-center pt-2 border-t border-zinc-200">
                                        <label class="text-xs font-medium flex items-center gap-1.5"><input type="checkbox" id="green2-chk" onchange="updateCompletionBadge(); saveLocalOnly();"> Financials</label>
                                        <input type="file" id="green2-file" multiple onchange="handleItemFiles('green2', this)" class="hidden">
                                        <div class="flex gap-1">
                                            <button type="button" onclick="openUploadPanel('green2')" class="bg-white hover:bg-zinc-100 text-zinc-700 text-[11px] py-1 px-2.5 rounded border border-zinc-300">📎 Browse</button>
                                            <button type="button" onclick="clearUploadField('green2')" class="bg-red-50 hover:bg-red-100 text-red-600 text-[11px] font-bold px-1.5 py-1 rounded border border-red-200 transition" title="Clear uploaded files">🗑️</button>
                                        </div>
                                    </div>
                                    <div id="files-green2" class="flex flex-wrap gap-1.5 mt-1 text-[11px] text-zinc-600">No files</div>
                                    <input type="text" id="green2-note" placeholder="Notes..." oninput="syncNotes('green2', this.value)" class="w-full bg-white border border-zinc-300 rounded px-2 py-1 text-xs">

                                    <div class="flex justify-between items-center pt-2 border-t border-zinc-200">
                                        <label class="text-xs font-medium flex items-center gap-1.5"><input type="checkbox" id="green3-chk" onchange="updateCompletionBadge(); saveLocalOnly();"> Payment Types</label>
                                        <input type="file" id="green3-file" multiple onchange="handleItemFiles('green3', this)" class="hidden">
                                        <div class="flex gap-1">
                                            <button type="button" onclick="openUploadPanel('green3')" class="bg-white hover:bg-zinc-100 text-zinc-700 text-[11px] py-1 px-2.5 rounded border border-zinc-300">📎 Browse</button>
                                            <button type="button" onclick="clearUploadField('green3')" class="bg-red-50 hover:bg-red-100 text-red-600 text-[11px] font-bold px-1.5 py-1 rounded border border-red-200 transition" title="Clear uploaded files">🗑️</button>
                                        </div>
                                    </div>
                                    <div id="files-green3" class="flex flex-wrap gap-1.5 mt-1 text-[11px] text-zinc-600">No files</div>
                                    <input type="text" id="green3-note" placeholder="Notes..." oninput="syncNotes('green3', this.value)" class="w-full bg-white border border-zinc-300 rounded px-2 py-1 text-xs">
                                </div>
                            </div>

                            <!-- 4. Moneris Reports -->
                            <div class="bg-white p-3 rounded-xl border border-zinc-200 space-y-2.5 shadow-xs">
                                <span class="text-xs font-bold text-emerald-800 block">Moneris Reports</span>
                                <div class="pl-3 space-y-2.5 bg-zinc-50 p-2.5 rounded-lg border border-zinc-200">
                                    <div class="flex justify-between items-center">
                                        <label class="text-xs font-medium flex items-center gap-1.5"><input type="checkbox" id="mon1-chk" onchange="updateCompletionBadge(); saveLocalOnly();"> Counter Sales</label>
                                        <input type="file" id="mon1-file" multiple onchange="handleItemFiles('mon1', this)" class="hidden">
                                        <div class="flex gap-1">
                                            <button type="button" onclick="openUploadPanel('mon1')" class="bg-white hover:bg-zinc-100 text-zinc-700 text-[11px] py-1 px-2.5 rounded border border-zinc-300">📎 Browse</button>
                                            <button type="button" onclick="clearUploadField('mon1')" class="bg-red-50 hover:bg-red-100 text-red-600 text-[11px] font-bold px-1.5 py-1 rounded border border-red-200 transition" title="Clear uploaded files">🗑️</button>
                                        </div>
                                    </div>
                                    <div id="files-mon1" class="flex flex-wrap gap-1.5 mt-1 text-[11px] text-zinc-600">No files</div>
                                    <input type="text" id="mon1-note" placeholder="Notes..." oninput="syncNotes('mon1', this.value)" class="w-full bg-white border border-zinc-300 rounded px-2 py-1 text-xs">

                                    <div class="flex justify-between items-center pt-2 border-t border-zinc-200">
                                        <label class="text-xs font-medium flex items-center gap-1.5"><input type="checkbox" id="mon2-chk" onchange="updateCompletionBadge(); saveLocalOnly();"> Online Sales</label>
                                        <input type="file" id="mon2-file" multiple onchange="handleItemFiles('mon2', this)" class="hidden">
                                        <div class="flex gap-1">
                                            <button type="button" onclick="openUploadPanel('mon2')" class="bg-white hover:bg-zinc-100 text-zinc-700 text-[11px] py-1 px-2.5 rounded border border-zinc-300">📎 Browse</button>
                                            <button type="button" onclick="clearUploadField('mon2')" class="bg-red-50 hover:bg-red-100 text-red-600 text-[11px] font-bold px-1.5 py-1 rounded border border-red-200 transition" title="Clear uploaded files">🗑️</button>
                                        </div>
                                    </div>
                                    <div id="files-mon2" class="flex flex-wrap gap-1.5 mt-1 text-[11px] text-zinc-600">No files</div>
                                    <input type="text" id="mon2-note" placeholder="Notes..." oninput="syncNotes('mon2', this.value)" class="w-full bg-white border border-zinc-300 rounded px-2 py-1 text-xs">
                                </div>
                            </div>

                            <!-- 5. OCS Documents -->
                            <div class="bg-white p-3 rounded-xl border border-zinc-200 space-y-2.5 shadow-xs">
                                <span class="text-xs font-bold text-emerald-800 block">OCS Documents</span>
                                <div class="pl-3 space-y-2.5 bg-zinc-50 p-2.5 rounded-lg border border-zinc-200">
                                    <div class="flex justify-between items-center">
                                        <label class="text-xs font-medium flex items-center gap-1.5"><input type="checkbox" id="ocs1-chk" onchange="updateCompletionBadge(); saveLocalOnly();"> Account Statement</label>
                                        <input type="file" id="ocs1-file" multiple onchange="handleItemFiles('ocs1', this)" class="hidden">
                                        <div class="flex gap-1">
                                            <button type="button" onclick="openUploadPanel('ocs1')" class="bg-white hover:bg-zinc-100 text-zinc-700 text-[11px] py-1 px-2.5 rounded border border-zinc-300">📎 Browse</button>
                                            <button type="button" onclick="clearUploadField('ocs1')" class="bg-red-50 hover:bg-red-100 text-red-600 text-[11px] font-bold px-1.5 py-1 rounded border border-red-200 transition" title="Clear uploaded files">🗑️</button>
                                        </div>
                                    </div>
                                    <div id="files-ocs1" class="flex flex-wrap gap-1.5 mt-1 text-[11px] text-zinc-600">No files</div>
                                    <input type="text" id="ocs1-note" placeholder="Notes..." oninput="syncNotes('ocs1', this.value)" class="w-full bg-white border border-zinc-300 rounded px-2 py-1 text-xs">

                                    <div class="flex justify-between items-center pt-2 border-t border-zinc-200">
                                        <label class="text-xs font-medium flex items-center gap-1.5"><input type="checkbox" id="ocs2-chk" onchange="updateCompletionBadge(); saveLocalOnly();"> Invoices</label>
                                        <input type="file" id="ocs2-file" multiple onchange="handleItemFiles('ocs2', this)" class="hidden">
                                        <div class="flex gap-1">
                                            <button type="button" onclick="openUploadPanel('ocs2')" class="bg-white hover:bg-zinc-100 text-zinc-700 text-[11px] py-1 px-2.5 rounded border border-zinc-300">📎 Browse</button>
                                            <button type="button" onclick="clearUploadField('ocs2')" class="bg-red-50 hover:bg-red-100 text-red-600 text-[11px] font-bold px-1.5 py-1 rounded border border-red-200 transition" title="Clear uploaded files">🗑️</button>
                                        </div>
                                    </div>
                                    <div id="files-ocs2" class="flex flex-wrap gap-1.5 mt-1 text-[11px] text-zinc-600">No files</div>
                                    <input type="text" id="ocs2-note" placeholder="Notes..." oninput="syncNotes('ocs2', this.value)" class="w-full bg-white border border-zinc-300 rounded px-2 py-1 text-xs">
                                </div>
                            </div>

                            <!-- 6. Moneris Chargebacks -->
                            <div class="bg-white p-3 rounded-xl border border-zinc-200 space-y-2 shadow-xs">
                                <span class="text-xs font-bold text-emerald-800 block">Moneris Chargebacks</span>
                                <div class="grid grid-cols-1 md:grid-cols-2 gap-2">
                                    <input type="text" id="chargebacks-text" value="Nothing to report." oninput="syncNotes('chargebacks', this.value)" class="w-full border border-zinc-300 rounded-lg px-3 py-2 text-xs">
                                    <div class="flex flex-col gap-1">
                                        <input type="file" id="chargebacks-file" multiple onchange="handleItemFiles('chargebacks', this)" class="hidden">
                                        <div class="flex gap-1">
                                            <button type="button" onclick="openUploadPanel('chargebacks')" class="flex-1 bg-zinc-100 hover:bg-zinc-200 text-zinc-700 text-xs font-medium py-1.5 px-3 rounded-lg border border-zinc-300 transition truncate">📎 Browse & Upload</button>
                                            <button type="button" onclick="clearUploadField('chargebacks')" class="bg-red-50 hover:bg-red-100 text-red-600 text-xs font-bold px-2 py-1.5 rounded-lg border border-red-200 transition" title="Clear uploaded files">🗑️</button>
                                        </div>
                                    </div>
                                </div>
                                <div id="files-chargebacks" class="flex flex-wrap gap-1.5 mt-1 text-[11px] text-zinc-600 bg-zinc-50 p-1.5 rounded border border-zinc-200">No files uploaded</div>
                            </div>

                            <!-- 7. Gas Cards -->
                            <div class="bg-white p-3 rounded-xl border border-zinc-200 space-y-2.5 shadow-xs">
                                <span class="text-xs font-bold text-emerald-800 block">Gas Cards</span>
                                <div class="pl-3 space-y-2.5 bg-zinc-50 p-2.5 rounded-lg border border-zinc-200">
                                    <div class="flex justify-between items-center">
                                        <label class="text-xs font-medium flex items-center gap-1.5"><input type="checkbox" id="gas1-chk" onchange="updateCompletionBadge(); saveLocalOnly();"> Details Report</label>
                                        <input type="file" id="gas1-file" multiple onchange="handleItemFiles('gas1', this)" class="hidden">
                                        <div class="flex gap-1">
                                            <button type="button" onclick="openUploadPanel('gas1')" class="bg-white hover:bg-zinc-100 text-zinc-700 text-[11px] py-1 px-2.5 rounded border border-zinc-300">📎 Browse</button>
                                            <button type="button" onclick="clearUploadField('gas1')" class="bg-red-50 hover:bg-red-100 text-red-600 text-[11px] font-bold px-1.5 py-1 rounded border border-red-200 transition" title="Clear uploaded files">🗑️</button>
                                        </div>
                                    </div>
                                    <div id="files-gas1" class="flex flex-wrap gap-1.5 mt-1 text-[11px] text-zinc-600">No files</div>
                                    <input type="text" id="gas1-note" placeholder="Notes..." oninput="syncNotes('gas1', this.value)" class="w-full bg-white border border-zinc-300 rounded px-2 py-1 text-xs">

                                    <div class="flex justify-between items-center pt-2 border-t border-zinc-200">
                                        <label class="text-xs font-medium flex items-center gap-1.5"><input type="checkbox" id="gas2-chk" onchange="updateCompletionBadge(); saveLocalOnly();"> Summary Report</label>
                                        <input type="file" id="gas2-file" multiple onchange="handleItemFiles('gas2', this)" class="hidden">
                                        <div class="flex gap-1">
                                            <button type="button" onclick="openUploadPanel('gas2')" class="bg-white hover:bg-zinc-100 text-zinc-700 text-[11px] py-1 px-2.5 rounded border border-zinc-300">📎 Browse</button>
                                            <button type="button" onclick="clearUploadField('gas2')" class="bg-red-50 hover:bg-red-100 text-red-600 text-[11px] font-bold px-1.5 py-1 rounded border border-red-200 transition" title="Clear uploaded files">🗑️</button>
                                        </div>
                                    </div>
                                    <div id="files-gas2" class="flex flex-wrap gap-1.5 mt-1 text-[11px] text-zinc-600">No files</div>
                                    <input type="text" id="gas2-note" placeholder="Notes..." oninput="syncNotes('gas2', this.value)" class="w-full bg-white border border-zinc-300 rounded px-2 py-1 text-xs">
                                </div>
                            </div>

                        </div>
                    </div>
                </div>

            </div>

            <!-- ================= ACCOUNTANT REVIEW PORTAL ================= -->
            <div class="bg-emerald-950 text-white rounded-2xl p-6 border border-emerald-800 shadow-lg">
                <div class="flex flex-col md:flex-row justify-between items-start md:items-center mb-4 pb-3 border-b border-emerald-800 gap-3">
                    <div>
                        <h3 class="text-lg font-bold text-emerald-300">Accountant Review Portal & Monthly Record Viewer</h3>
                        <p class="text-emerald-200/80 text-xs mt-0.5">Accountant view for <span id="portal-store-label" class="text-white font-semibold">Oz 1</span> – <span id="portal-period-label" class="text-white font-semibold">October 2026</span></p>
                    </div>
                    <div class="flex items-center gap-2">
                        <button onclick="downloadMasterZip()" class="bg-emerald-700 hover:bg-emerald-600 text-white text-xs font-semibold px-3.5 py-2 rounded-xl border border-emerald-500 transition">
                            📦 Download All as ZIP
                        </button>
                        <button onclick="window.print()" class="bg-emerald-900 hover:bg-emerald-800 text-emerald-100 text-xs font-semibold px-3.5 py-2 rounded-xl border border-emerald-700 transition">
                            Print / Save PDF Summary
                        </button>
                    </div>
                </div>

                <div class="grid grid-cols-1 md:grid-cols-3 gap-4 mb-6">
                    <div class="bg-emerald-900/60 p-4 rounded-xl border border-emerald-800">
                        <span class="text-xs font-semibold text-emerald-300 block mb-1">Cash On Hand</span>
                        <span id="portal-cooh-val" class="text-base font-bold text-white">Not entered</span>
                    </div>
                    <div class="bg-emerald-900/60 p-4 rounded-xl border border-emerald-800">
                        <span class="text-xs font-semibold text-emerald-300 block mb-1">Inventory Formatted Report</span>
                        <span id="portal-inv-val" class="text-xs font-medium text-emerald-200">No CSV report uploaded</span>
                    </div>
                    <div class="bg-emerald-900/60 p-4 rounded-xl border border-emerald-800">
                        <span class="text-xs font-semibold text-emerald-300 block mb-1">Checklist Status</span>
                        <span id="portal-chk-val" class="text-xs font-bold text-amber-300">0 / 10 Items Checked</span>
                    </div>
                </div>

                <div class="bg-emerald-900/40 rounded-xl border border-emerald-800 overflow-hidden">
                    <table class="w-full text-left border-collapse text-xs">
                        <thead class="bg-emerald-900 text-emerald-300 border-b border-emerald-800">
                            <tr>
                                <th class="p-3 font-semibold">Document / Folder Category</th>
                                <th class="p-3 font-semibold">Status</th>
                                <th class="p-3 font-semibold">Document Date</th>
                                <th class="p-3 font-semibold">Attached File Names</th>
                                <th class="p-3 font-semibold">Notes & Details</th>
                            </tr>
                        </thead>
                        <tbody id="portal-table-body" class="divide-y divide-emerald-800/60 text-emerald-100">
                            <tr class="hover:bg-emerald-900/30"><td class="p-3 font-medium text-white">Cash On Hand - End of Day</td><td class="p-3"><span class="bg-emerald-800 text-emerald-200 px-2 py-0.5 rounded text-[11px]">Entered ✓</span></td><td class="p-3 text-emerald-300/50 italic" id="portal-date-cooh">No date</td><td class="p-3" id="table-files-cooh">No files</td><td class="p-3 text-emerald-100" id="note-cooh">Not entered</td></tr>
                            <tr class="hover:bg-emerald-900/30"><td class="p-3 font-medium text-white">Brinks Book (Photos highlighted)</td><td class="p-3"><span class="bg-emerald-900 text-emerald-300/60 px-2 py-0.5 rounded text-[11px]" id="status-brinks">Pending</span></td><td class="p-3 text-emerald-300/50 italic" id="portal-date-brinks">No date</td><td class="p-3" id="table-files-brinks">No files</td><td class="p-3 text-emerald-100" id="note-brinks">No notes</td></tr>
                            <tr class="hover:bg-emerald-900/30"><td class="p-3 font-medium text-white">Greenline: Inventory Summary</td><td class="p-3"><span class="bg-emerald-900 text-emerald-300/60 px-2 py-0.5 rounded text-[11px]" id="status-green1">Pending</span></td><td class="p-3 text-emerald-300/50 italic">No date</td><td class="p-3" id="table-files-green1">No files</td><td class="p-3 text-emerald-100" id="note-green1">No notes</td></tr>
                            <tr class="hover:bg-emerald-900/30"><td class="p-3 font-medium text-white">Greenline: Financials</td><td class="p-3"><span class="bg-emerald-900 text-emerald-300/60 px-2 py-0.5 rounded text-[11px]" id="status-green2">Pending</span></td><td class="p-3 text-emerald-300/50 italic">No date</td><td class="p-3" id="table-files-green2">No files</td><td class="p-3 text-emerald-100" id="note-green2">No notes</td></tr>
                            <tr class="hover:bg-emerald-900/30"><td class="p-3 font-medium text-white">Greenline: Payment Types</td><td class="p-3"><span class="bg-emerald-900 text-emerald-300/60 px-2 py-0.5 rounded text-[11px]" id="status-green3">Pending</span></td><td class="p-3 text-emerald-300/50 italic">No date</td><td class="p-3" id="table-files-green3">No files</td><td class="p-3 text-emerald-100" id="note-green3">No notes</td></tr>
                            <tr class="hover:bg-emerald-900/30"><td class="p-3 font-medium text-white">Moneris: Counter Sales Statement</td><td class="p-3"><span class="bg-emerald-900 text-emerald-300/60 px-2 py-0.5 rounded text-[11px]" id="status-mon1">Pending</span></td><td class="p-3 text-emerald-300/50 italic">No date</td><td class="p-3" id="table-files-mon1">No files</td><td class="p-3 text-emerald-100" id="note-mon1">No notes</td></tr>
                            <tr class="hover:bg-emerald-900/30"><td class="p-3 font-medium text-white">Moneris: Online Sales Statement</td><td class="p-3"><span class="bg-emerald-900 text-emerald-300/60 px-2 py-0.5 rounded text-[11px]" id="status-mon2">Pending</span></td><td class="p-3 text-emerald-300/50 italic">No date</td><td class="p-3" id="table-files-mon2">No files</td><td class="p-3 text-emerald-100" id="note-mon2">No notes</td></tr>
                            <tr class="hover:bg-emerald-900/30"><td class="p-3 font-medium text-white">OCS: Account Statement</td><td class="p-3"><span class="bg-emerald-900 text-emerald-300/60 px-2 py-0.5 rounded text-[11px]" id="status-ocs1">Pending</span></td><td class="p-3 text-emerald-300/50 italic">No date</td><td class="p-3" id="table-files-ocs1">No files</td><td class="p-3 text-emerald-100" id="note-ocs1">No notes</td></tr>
                            <tr class="hover:bg-emerald-900/30"><td class="p-3 font-medium text-white">OCS: Invoices (Multiple)</td><td class="p-3"><span class="bg-emerald-900 text-emerald-300/60 px-2 py-0.5 rounded text-[11px]" id="status-ocs2">Pending</span></td><td class="p-3 text-emerald-300/50 italic">No date</td><td class="p-3" id="table-files-ocs2">No files</td><td class="p-3 text-emerald-100" id="note-ocs2">No notes</td></tr>
                            <tr class="hover:bg-emerald-900/30"><td class="p-3 font-medium text-white">Moneris Chargebacks</td><td class="p-3"><span class="bg-emerald-800 text-emerald-200 px-2 py-0.5 rounded text-[11px]">Entered ✓</span></td><td class="p-3 text-emerald-300/50 italic">No date</td><td class="p-3" id="table-files-chargebacks">No files</td><td class="p-3 text-emerald-100" id="note-chargebacks">Nothing to report.</td></tr>
                            <tr class="hover:bg-emerald-900/30"><td class="p-3 font-medium text-white">Gas Cards: Details Report</td><td class="p-3"><span class="bg-emerald-900 text-emerald-300/60 px-2 py-0.5 rounded text-[11px]" id="status-gas1">Pending</span></td><td class="p-3 text-emerald-300/50 italic">No date</td><td class="p-3" id="table-files-gas1">No files</td><td class="p-3 text-emerald-100" id="note-gas1">No notes</td></tr>
                            <tr class="hover:bg-emerald-900/30"><td class="p-3 font-medium text-white">Gas Cards: Summary Report</td><td class="p-3"><span class="bg-emerald-900 text-emerald-300/60 px-2 py-0.5 rounded text-[11px]" id="status-gas2">Pending</span></td><td class="p-3 text-emerald-300/50 italic">No date</td><td class="p-3" id="table-files-gas2">No files</td><td class="p-3 text-emerald-100" id="note-gas2">No notes</td></tr>
                        </tbody>
                    </table>
                </div>

                <!-- Monthly Financial Summary Report (Executive Owner View) -->
                <div class="mt-6 bg-emerald-900/40 rounded-xl border border-emerald-800 p-5">
                    <h4 class="text-base font-bold text-emerald-300 mb-4">Monthly Financial Summary & Owner Expense Report</h4>

                    <div class="grid grid-cols-1 md:grid-cols-2 lg:grid-cols-4 gap-4 mb-6">
                        <div class="bg-emerald-800/50 p-3 rounded-lg">
                            <span class="text-xs text-emerald-300 block">Total Sales (Inventory Sold)</span>
                            <span id="fin-total-sales" class="text-lg font-bold text-white">$0.00</span>
                        </div>
                        <div class="bg-emerald-800/50 p-3 rounded-lg">
                            <span class="text-xs text-emerald-300 block">Cost of Goods Sold (COGS)</span>
                            <span id="fin-cogs" class="text-lg font-bold text-emerald-200">$0.00</span>
                            <span id="fin-gross-margin" class="text-[11px] text-emerald-200 block">Gross Margin: 0%</span>
                        </div>
                        <div class="bg-emerald-800/50 p-3 rounded-lg">
                            <span class="text-xs text-emerald-300 block">Net Operating Income (NOI)</span>
                            <span id="fin-net" class="text-lg font-bold text-white">$0.00</span>
                        </div>
                        <div class="bg-emerald-800/50 p-3 rounded-lg">
                            <span class="text-xs text-emerald-300 block">Operating Margin</span>
                            <span id="fin-net-margin" class="text-lg font-bold text-amber-300">0.0%</span>
                        </div>
                    </div>

                    <h5 class="text-sm font-bold text-emerald-300 mb-3">Operating Expenses (Editable Fields)</h5>
                    <div class="grid grid-cols-1 sm:grid-cols-2 lg:grid-cols-5 gap-3 mb-4">
                        <div>
                            <label class="text-[11px] text-emerald-200 block">Store Rent & Lease</label>
                            <input type="text" id="exp-rent" value="2000.00" oninput="updateFinancialSummary(); saveLocalOnly();" class="w-full bg-emerald-950 text-white border border-emerald-700 rounded px-2 py-1 text-xs">
                        </div>
                        <div>
                            <label class="text-[11px] text-emerald-200 block">Employee Salaries</label>
                            <input type="text" id="exp-salaries" value="1950.00" oninput="updateFinancialSummary(); saveLocalOnly();" class="w-full bg-emerald-950 text-white border border-emerald-700 rounded px-2 py-1 text-xs">
                        </div>
                        <div>
                            <label class="text-[11px] text-emerald-200 block">Utilities</label>
                            <input type="text" id="exp-utilities" value="1200.00" oninput="updateFinancialSummary(); saveLocalOnly();" class="w-full bg-emerald-950 text-white border border-emerald-700 rounded px-2 py-1 text-xs">
                        </div>
                        <div>
                            <label class="text-[11px] text-emerald-200 block">Advertising</label>
                            <input type="text" id="exp-advertising" value="1000.00" oninput="updateFinancialSummary(); saveLocalOnly();" class="w-full bg-emerald-950 text-white border border-emerald-700 rounded px-2 py-1 text-xs">
                        </div>
                        <div>
                            <label class="text-[11px] text-emerald-200 block">Marketing</label>
                            <input type="text" id="exp-marketing" value="800.00" oninput="updateFinancialSummary(); saveLocalOnly();" class="w-full bg-emerald-950 text-white border border-emerald-700 rounded px-2 py-1 text-xs">
                        </div>
                        <div>
                            <label class="text-[11px] text-emerald-200 block">Data Deals & Analytics</label>
                            <input type="text" id="exp-datadeals" value="1500.00" oninput="updateFinancialSummary(); saveLocalOnly();" class="w-full bg-emerald-950 text-white border border-emerald-700 rounded px-2 py-1 text-xs">
                        </div>
                        <div>
                            <label class="text-[11px] text-emerald-200 block">Insurance</label>
                            <input type="text" id="exp-insurance" value="500.00" oninput="updateFinancialSummary(); saveLocalOnly();" class="w-full bg-emerald-950 text-white border border-emerald-700 rounded px-2 py-1 text-xs">
                        </div>
                        <div>
                            <label class="text-[11px] text-emerald-200 block">Garbage Disposal</label>
                            <input type="text" id="exp-garbage" value="150.00" oninput="updateFinancialSummary(); saveLocalOnly();" class="w-full bg-emerald-950 text-white border border-emerald-700 rounded px-2 py-1 text-xs">
                        </div>
                        <div>
                            <label class="text-[11px] text-emerald-200 block">Security & POS Software</label>
                            <input type="text" id="exp-security" value="350.00" oninput="updateFinancialSummary(); saveLocalOnly();" class="w-full bg-emerald-950 text-white border border-emerald-700 rounded px-2 py-1 text-xs">
                        </div>
                        <div>
                            <label class="text-[11px] text-emerald-200 block">Licenses & Compliance</label>
                            <input type="text" id="exp-licenses" value="300.00" oninput="updateFinancialSummary(); saveLocalOnly();" class="w-full bg-emerald-950 text-white border border-emerald-700 rounded px-2 py-1 text-xs">
                        </div>
                    </div>

                    <div class="bg-emerald-800/30 p-3 rounded-lg mb-3 flex justify-between items-center">
                        <span class="text-xs text-emerald-300 font-bold">Total Operating Expenses (OpEx):</span>
                        <span id="fin-total-expenses" class="text-sm font-bold text-white">$0.00</span>
                    </div>

                    <div class="text-xs text-emerald-200/70">
                        <p>Total Sales and Cost of Goods Sold (COGS) are auto-calculated from your uploaded Greenline inventory CSV. Operating Expenses are fully editable above and saved when you press <b>Save All to Cloud</b>.</p>
                    </div>
                </div>
            </div>

        </div>
    </div>

    <script>
        const ACCOUNTANT_ORDER = ["Beverages", "Dried Flower", "Edibles", "Extracts Ingested", "Extracts Inhaled", "Infused Pre-Rolls ", "Milled Flower", "Pre-Rolls", "Seeds", "Topicals", "Vapes - Disposables", "Vapes - 510 Threads", "Vapes - Pods", "Accessories", "Apparel / Swag", "Gift Cards", "Other Income", "Raffles", "System Fees", "Uniforms"];

        let isStoreManager = REPLACE_IS_STORE_MANAGER;
        let currentUserRole = "REPLACE_CURRENT_ROLE";
        let currentUserName = "REPLACE_CURRENT_NAME";
        let currentUsername = "REPLACE_CURRENT_USERNAME";

        let stores = (isStoreManager || currentUserRole === 'store') ? [currentUserName] : ["Oz 1", "Oz 2", "Oz 3", "Oz 4", "Oz 5", "Oz 6", "Oz 7"];
        let activeStore = "REPLACE_ACTIVE_STORE";
        if (!stores.includes(activeStore)) activeStore = stores[0];
        let selectedMonth = "REPLACE_MONTH";
        let selectedYear = "REPLACE_YEAR";

        let storedFilesMap = REPLACE_FILES_JSON;
        let persistentStoreData = REPLACE_STORE_DATA_JSON;
        let storeInventoryData = {};

        function esc(s) {
            return String(s === undefined || s === null ? '' : s)
                .replace(/&/g, '&amp;').replace(/</g, '&lt;').replace(/>/g, '&gt;')
                .replace(/"/g, '&quot;').replace(/'/g, '&#39;');
        }

        function fmtSize(n) {
            n = Number(n) || 0;
            if (n >= 1048576) return (n / 1048576).toFixed(1) + ' MB';
            if (n >= 1024) return Math.round(n / 1024) + ' KB';
            return n ? n + ' B' : '';
        }

        function ctxQS(extra) {
            return '?auth_user=' + encodeURIComponent(currentUsername) +
                '&store=' + encodeURIComponent(activeStore) +
                '&month=' + encodeURIComponent(selectedMonth) +
                '&year=' + encodeURIComponent(selectedYear) + (extra || '');
        }

        function navigateContext(store, month, year) {
            saveLocalOnly();
            navigateParent('?auth_user=' + encodeURIComponent(currentUsername) +
                '&store=' + encodeURIComponent(store) +
                '&month=' + encodeURIComponent(month) +
                '&year=' + encodeURIComponent(year));
        }

        function openUploadPanel(cat) {
            saveLocalOnly();
            navigateParent(ctxQS('&action=upload_panel&cat=' + encodeURIComponent(cat)));
        }

        function navigateParent(queryString) {
            try {
                window.parent.location.search = queryString;
            } catch (e) {
                window.location.href = queryString;
            }
        }

        function initApp() {
            document.getElementById('select-month').value = selectedMonth;
            document.getElementById('select-year').value = selectedYear;
            document.getElementById('active-period-label').textContent = `${selectedMonth} ${selectedYear}`;
            document.getElementById('portal-period-label').textContent = `${selectedMonth} ${selectedYear}`;
            renderTabs();
            loadStorePeriodData();
            setupCSVParser();
            document.getElementById('workspace-store-name').textContent = activeStore;
            document.getElementById('portal-store-label').textContent = activeStore;
        }

        function renderTabs() {
            const storeTabsEl = document.getElementById('store-tabs');
            if(!storeTabsEl) return;
            storeTabsEl.innerHTML = stores.map(store => `
                <button onclick="switchStore(this.dataset.store)" data-store="${esc(store)}" class="px-3.5 py-1.5 rounded-xl text-xs font-bold whitespace-nowrap transition-all ${activeStore === store ? 'bg-emerald-700 text-white shadow-md' : 'bg-white hover:bg-zinc-100 text-zinc-700 border border-zinc-300'}">
                    ${esc(store)}
                </button>
            `).join('');
        }

        function switchStore(store) {
            if (store === activeStore) return;
            navigateContext(store, selectedMonth, selectedYear);
        }

        function updatePeriod() {
            const m = document.getElementById('select-month').value;
            const y = document.getElementById('select-year').value;
            if (m === selectedMonth && y === selectedYear) return;
            navigateContext(activeStore, m, y);
        }

        function getSPK() {
            return `${activeStore}_${selectedMonth}_${selectedYear}`;
        }

        function loadStorePeriodData() {
            const spk = getSPK();

            document.getElementById('cooh-text').value = "";
            document.getElementById('cooh-date').value = "";
            document.getElementById('brinks-date').value = "";
            document.getElementById('chargebacks-text').value = "Nothing to report.";

            const checkboxes = ['brinks', 'green1', 'green2', 'green3', 'mon1', 'mon2', 'ocs1', 'ocs2', 'gas1', 'gas2'];
            checkboxes.forEach(k => {
                const chk = document.getElementById(`${k}-chk`);
                if(chk) chk.checked = false;
                const note = document.getElementById(`${k}-note`);
                if(note) note.value = "";
                syncNotes(k, "", false);
            });

            let data = persistentStoreData[spk];
            try {
                const local = localStorage.getItem('oz_data_' + spk);
                if (local) {
                    const parsedLocal = JSON.parse(local);
                    if (!data || parsedLocal._updated > (data._updated || 0)) {
                        data = parsedLocal;
                    }
                }
            } catch(e) {}

            if (data) {
                if(data.cooh_text !== undefined) {
                    document.getElementById('cooh-text').value = data.cooh_text;
                    syncNotes('cooh', data.cooh_text, false);
                }
                if(data.cooh_date) document.getElementById('cooh-date').value = data.cooh_date;
                if(data.chargebacks_text) {
                    document.getElementById('chargebacks-text').value = data.chargebacks_text;
                    syncNotes('chargebacks', data.chargebacks_text, false);
                }
                if(data.checklist_json) {
                    const chkData = data.checklist_json;
                    if(chkData.checkboxes) {
                        Object.keys(chkData.checkboxes).forEach(k => {
                            const chk = document.getElementById(`${k}-chk`);
                            if(chk) chk.checked = chkData.checkboxes[k];
                        });
                    }
                    if(chkData.notes) {
                        Object.keys(chkData.notes).forEach(k => {
                            const nEl = document.getElementById(`${k}-note`);
                            if(nEl) {
                                nEl.value = chkData.notes[k];
                                syncNotes(k, chkData.notes[k], false);
                            }
                        });
                    }
                    if(chkData.dates) {
                        if(chkData.dates.brinks) document.getElementById('brinks-date').value = chkData.dates.brinks;
                    }
                    if (chkData.expenses) {
                        const exp = chkData.expenses;
                        if (exp.rent !== undefined) document.getElementById('exp-rent').value = exp.rent;
                        if (exp.salaries !== undefined) document.getElementById('exp-salaries').value = exp.salaries;
                        if (exp.utilities !== undefined) document.getElementById('exp-utilities').value = exp.utilities;
                        if (exp.advertising !== undefined) document.getElementById('exp-advertising').value = exp.advertising;
                        if (exp.marketing !== undefined) document.getElementById('exp-marketing').value = exp.marketing;
                        if (exp.datadeals !== undefined) document.getElementById('exp-datadeals').value = exp.datadeals;
                        if (exp.insurance !== undefined) document.getElementById('exp-insurance').value = exp.insurance;
                        if (exp.garbage !== undefined) document.getElementById('exp-garbage').value = exp.garbage;
                        if (exp.security !== undefined) document.getElementById('exp-security').value = exp.security;
                        if (exp.licenses !== undefined) document.getElementById('exp-licenses').value = exp.licenses;
                    }
                }
                const invObj = data.inventory_data || data.inventory_json;
                if(invObj) {
                    storeInventoryData[spk] = {
                        headers: invObj.headers,
                        rows: invObj.rows,
                        fileName: data.inventory_filename || 'Uploaded_Inventory.csv'
                    };
                } else {
                    delete storeInventoryData[spk];
                }
            }

            renderInventoryTable();
            renderStoredFiles();
            updateCompletionBadge();
            updateFinancialSummary();
        }

        function setupCSVParser() {
            const dropZone = document.getElementById('drop-zone');
            const fileInput = document.getElementById('file-input');
            if(!dropZone || !fileInput) return;

            dropZone.onclick = () => fileInput.click();
            dropZone.ondragover = (e) => { e.preventDefault(); dropZone.classList.add('border-emerald-600', 'bg-emerald-50/50'); };
            dropZone.ondragleave = () => dropZone.classList.remove('border-emerald-600', 'bg-emerald-50/50');
            dropZone.ondrop = (e) => {
                e.preventDefault();
                dropZone.classList.remove('border-emerald-600', 'bg-emerald-50/50');
                if (e.dataTransfer.files.length) handleCSVFile(e.dataTransfer.files[0]);
            };
            fileInput.onchange = (e) => {
                if (e.target.files.length) handleCSVFile(e.target.files[0]);
                fileInput.value = '';
            };
        }

        function handleCSVFile(file) {
            Papa.parse(file, {
                header: true,
                skipEmptyLines: true,
                complete: function(results) {
                    processParsedCSV(results.meta.fields, results.data, file.name);
                }
            });
        }

        function processParsedCSV(fields, data, fileName) {
            const parentRows = data.filter(row => {
                const cat = row['Category'] || '';
                return !cat.trim().startsWith('>');
            });

            const sortedRows = [];
            ACCOUNTANT_ORDER.forEach(targetCat => {
                const found = parentRows.find(r => (r['Category'] || '').trim() === targetCat.trim());
                if (found) {
                    sortedRows.push(found);
                } else {
                    const emptyRow = { Category: targetCat };
                    fields.forEach(f => { if (f !== 'Category') emptyRow[f] = 0.00; });
                    sortedRows.push(emptyRow);
                }
            });

            const spk = getSPK();
            storeInventoryData[spk] = { headers: fields, rows: sortedRows, fileName: fileName };

            const green1Chk = document.getElementById('green1-chk');
            if (green1Chk) green1Chk.checked = true;

            renderInventoryTable();
            showCSVFileInfo(fileName);
            renderStoredFiles();
            updateCompletionBadge();
            updateFinancialSummary();
            saveLocalOnly();
        }

        function renderInventoryTable() {
            const spk = getSPK();
            const inv = storeInventoryData[spk];
            const tableHeader = document.getElementById('table-header');
            const tableBody = document.getElementById('table-body');
            const invStatusBadge = document.getElementById('inv-status-badge');
            const actionArea = document.getElementById('action-area');
            const portalInvVal = document.getElementById('portal-inv-val');
            const infoBox = document.getElementById('csv-file-info');

            if (!inv) {
                tableHeader.innerHTML = `<th class="p-2 font-semibold">Category</th><th class="p-2 font-semibold">Opening</th><th class="p-2 font-semibold">Sold</th><th class="p-2 font-semibold">Closing</th>`;
                tableBody.innerHTML = `<tr><td colspan="4" class="p-4 text-center text-zinc-400">No report uploaded yet.</td></tr>`;
                actionArea.classList.add('hidden');
                invStatusBadge.textContent = "Not Loaded";
                invStatusBadge.className = "text-[11px] text-zinc-600 bg-white border border-zinc-300 px-2 py-0.5 rounded shadow-xs";
                portalInvVal.innerHTML = '<span class="text-emerald-200/60">No CSV report uploaded</span>';
                if(infoBox) infoBox.classList.add('hidden');
                updateFinancialSummary();
                return;
            }

            const headers = inv.headers;
            const rows = inv.rows;
            const displayHeaders = headers.includes('Opening') ? ['Category', 'Opening', 'Sold', 'Closing'] : [headers[0], headers[1], headers[2]];

            tableHeader.innerHTML = displayHeaders.map(h => `<th class="p-2 font-semibold">${esc(h)}</th>`).join('');
            tableBody.innerHTML = rows.map(row => `
                <tr class="hover:bg-zinc-50 transition-colors">
                    ${displayHeaders.map(h => `<td class="p-2 ${h === 'Category' ? 'font-medium text-emerald-800' : 'text-zinc-700'}">${row[h] !== undefined ? esc(row[h]) : ''}</td>`).join('')}
                </tr>
            `).join('');

            actionArea.classList.remove('hidden');
            invStatusBadge.textContent = `Loaded (${inv.fileName})`;
            invStatusBadge.className = "text-[11px] text-emerald-700 bg-emerald-50 border border-emerald-200 px-2 py-0.5 rounded";
            portalInvVal.innerHTML = `<span class="text-emerald-300 font-semibold">✓ Ready (${esc(inv.fileName)})</span>`;

            showCSVFileInfo(inv.fileName);

            document.getElementById('download-excel-btn').onclick = function() {
                const worksheetData = [headers, ...rows.map(row => headers.map(h => row[h]))];
                const worksheet = XLSX.utils.aoa_to_sheet(worksheetData);
                const workbook = XLSX.utils.book_new();
                XLSX.utils.book_append_sheet(workbook, worksheet, `${activeStore} Inventory`);
                XLSX.writeFile(workbook, `${activeStore}_${selectedMonth}_${selectedYear}_Formatted_Report.xlsx`);
            };
            updateFinancialSummary();
        }

        function showCSVFileInfo(fileName) {
            const infoBox = document.getElementById('csv-file-info');
            const nameSpan = document.getElementById('csv-file-name');
            if (infoBox && nameSpan) {
                nameSpan.textContent = fileName;
                infoBox.classList.remove('hidden');
            }
        }

        function clearCSVUpload() {
            if (confirm('Are you sure you want to delete this CSV upload?')) {
                const spk = getSPK();
                delete storeInventoryData[spk];
                const green1Chk = document.getElementById('green1-chk');
                if (green1Chk) green1Chk.checked = false;
                renderInventoryTable();
                renderStoredFiles();
                updateCompletionBadge();
                saveLocalOnly();
            }
        }

        function syncNotes(key, val, triggerSave = true) {
            const noteCell = document.getElementById(`note-${key}`);
            if (noteCell) {
                noteCell.textContent = val.trim() !== '' ? val : (key === 'chargebacks' ? 'Nothing to report.' : 'No notes');
            }
            if (key === 'cooh') {
                const coohValCell = document.getElementById('portal-cooh-val');
                if (coohValCell) coohValCell.textContent = val.trim() !== '' ? val : 'Not entered';
            }
            if (triggerSave) saveLocalOnly();
        }

        function buildPayload() {
            const spk = getSPK();
            const checkboxes = {
                'brinks': document.getElementById('brinks-chk')?.checked || false,
                'green1': document.getElementById('green1-chk')?.checked || false,
                'green2': document.getElementById('green2-chk')?.checked || false,
                'green3': document.getElementById('green3-chk')?.checked || false,
                'mon1': document.getElementById('mon1-chk')?.checked || false,
                'mon2': document.getElementById('mon2-chk')?.checked || false,
                'ocs1': document.getElementById('ocs1-chk')?.checked || false,
                'ocs2': document.getElementById('ocs2-chk')?.checked || false,
                'gas1': document.getElementById('gas1-chk')?.checked || false,
                'gas2': document.getElementById('gas2-chk')?.checked || false
            };

            const notes = {
                'cooh': document.getElementById('cooh-text')?.value || '',
                'brinks': document.getElementById('brinks-note')?.value || '',
                'green1': document.getElementById('green1-note')?.value || '',
                'green2': document.getElementById('green2-note')?.value || '',
                'green3': document.getElementById('green3-note')?.value || '',
                'mon1': document.getElementById('mon1-note')?.value || '',
                'mon2': document.getElementById('mon2-note')?.value || '',
                'ocs1': document.getElementById('ocs1-note')?.value || '',
                'ocs2': document.getElementById('ocs2-note')?.value || '',
                'chargebacks': document.getElementById('chargebacks-text')?.value || 'Nothing to report.',
                'gas1': document.getElementById('gas1-note')?.value || '',
                'gas2': document.getElementById('gas2-note')?.value || ''
            };

            const dates = {
                'cooh': document.getElementById('cooh-date')?.value || '',
                'brinks': document.getElementById('brinks-date')?.value || ''
            };

            const expenses = {
                rent: document.getElementById('exp-rent')?.value || '',
                salaries: document.getElementById('exp-salaries')?.value || '',
                utilities: document.getElementById('exp-utilities')?.value || '',
                advertising: document.getElementById('exp-advertising')?.value || '',
                marketing: document.getElementById('exp-marketing')?.value || '',
                datadeals: document.getElementById('exp-datadeals')?.value || '',
                insurance: document.getElementById('exp-insurance')?.value || '',
                garbage: document.getElementById('exp-garbage')?.value || '',
                security: document.getElementById('exp-security')?.value || '',
                licenses: document.getElementById('exp-licenses')?.value || ''
            };

            const checklistData = {
                checkboxes: checkboxes,
                notes: notes,
                dates: dates,
                expenses: expenses,
                files: {}
            };
            const inv = storeInventoryData[spk] || null;

            return {
                _updated: Date.now(),
                spk: spk,
                store: activeStore,
                period: selectedMonth + ' ' + selectedYear,
                cooh_text: notes['cooh'],
                cooh_date: dates['cooh'],
                chargebacks_text: notes['chargebacks'],
                checklist_json: checklistData,
                inventory_data: inv ? { headers: inv.headers, rows: inv.rows } : null,
                inventory_filename: inv ? inv.fileName : ""
            };
        }

        function saveLocalOnly() {
            const spk = getSPK();
            const payload = buildPayload();
            persistentStoreData[spk] = payload;
            try {
                localStorage.setItem('oz_data_' + spk, JSON.stringify(payload));
            } catch(e) {}
            const badge = document.getElementById('store-completion-badge');
            if (badge) {
                const checked = Object.values(payload.checklist_json.checkboxes).filter(Boolean).length;
                badge.textContent = `✓ Draft Saved | Checklist: ${checked}/10`;
            }
        }

        function saveAllData(isSilent = false) {
            const payload = buildPayload();
            const spk = payload.spk;
            persistentStoreData[spk] = payload;
            try {
                localStorage.setItem('oz_data_' + spk, JSON.stringify(payload));
            } catch(e) {}

            const payloadEncoded = encodeURIComponent(JSON.stringify(payload));
            if (payloadEncoded.length > 60000) {
                alert('This store has too much data to save in one go. Please remove unneeded rows from the inventory and try again.');
                return;
            }
            navigateParent(ctxQS('&action=save_all&payload=' + payloadEncoded));
        }

        function updateCompletionBadge() {
            const checkboxes = ['brinks', 'green1', 'green2', 'green3', 'mon1', 'mon2', 'ocs1', 'ocs2', 'gas1', 'gas2'];
            let checked = 0;
            checkboxes.forEach(key => {
                const chk = document.getElementById(`${key}-chk`);
                if (chk && chk.checked) {
                    checked++;
                    const statusEl = document.getElementById(`status-${key}`);
                    if(statusEl) {
                        statusEl.textContent = "Complete ✓";
                        statusEl.className = "bg-emerald-800 text-emerald-200 px-2 py-0.5 rounded text-[11px]";
                    }
                } else {
                    const statusEl = document.getElementById(`status-${key}`);
                    if(statusEl) {
                        statusEl.textContent = "Pending";
                        statusEl.className = "bg-emerald-900 text-emerald-300/60 px-2 py-0.5 rounded text-[11px]";
                    }
                }
            });

            const badge = document.getElementById('store-completion-badge');
            if (badge) badge.textContent = `Checklist: ${checked}/10 Completed`;

            const portalChk = document.getElementById('portal-chk-val');
            if (portalChk) portalChk.textContent = `${checked} / 10 Items Checked`;

            const coohDate = document.getElementById('cooh-date')?.value;
            const brinksDate = document.getElementById('brinks-date')?.value;
            if(document.getElementById('portal-date-cooh')) document.getElementById('portal-date-cooh').textContent = coohDate || 'No date';
            if(document.getElementById('portal-date-brinks')) document.getElementById('portal-date-brinks').textContent = brinksDate || 'No date';
        }

        function handleItemFiles(cat, inputElem) {
            if (inputElem) inputElem.value = '';
            openUploadPanel(cat);
        }

        function handleRemoveFile(cat, fileId) {
            if (confirm('Remove this file permanently?')) {
                saveLocalOnly();
                navigateParent(ctxQS('&action=delete_file&cat=' + encodeURIComponent(cat) + '&file_id=' + encodeURIComponent(fileId)));
            }
        }

        function clearUploadField(cat) {
            if (confirm('Clear all files for this category?')) {
                saveLocalOnly();
                navigateParent(ctxQS('&action=clear_files&cat=' + encodeURIComponent(cat)));
            }
        }

        function renderStoredFiles() {
            const spk = getSPK();
            const cats = ['cooh', 'brinks', 'green1', 'green2', 'green3', 'mon1', 'mon2', 'ocs1', 'ocs2', 'chargebacks', 'gas1', 'gas2'];

            cats.forEach(cat => {
                const box = document.getElementById(`files-${cat}`);
                const tableBox = document.getElementById(`table-files-${cat}`);
                if (!box) return;

                const fileList = (storedFilesMap[spk] && storedFilesMap[spk][cat]) ? storedFilesMap[spk][cat] : [];

                if (fileList.length > 0) {
                    const chip = (f, cls, xcls) => {
                        const size = fmtSize(f.size);
                        return `<span class="inline-block ${cls} px-1.5 py-0.5 rounded mr-1 mb-1 font-medium">📎 ${esc(f.filename)}${size ? ' <span class="opacity-60">(' + size + ')</span>' : ''}` +
                            ` <button type="button" data-cat="${esc(cat)}" data-id="${esc(f.id)}" onclick="handleRemoveFile(this.dataset.cat, this.dataset.id)" class="${xcls} font-bold ml-1">×</button></span>`;
                    };
                    box.innerHTML = fileList.map(f => chip(f, 'bg-emerald-100 text-emerald-800', 'text-red-600 hover:text-red-800')).join('');
                    if (tableBox) {
                        tableBox.innerHTML = fileList.map(f => chip(f, 'bg-emerald-800 text-emerald-200', 'text-red-400 hover:text-red-300')).join('');
                    }
                } else {
                    box.innerHTML = '<span class="text-zinc-400 italic">No files uploaded</span>';
                    if (tableBox) tableBox.innerHTML = '<span class="text-emerald-300/50 italic">No files attached</span>';
                }
            });
        }

        function downloadMasterZip() {
            // Save what is on screen first, then ask the server to build the real package
            // (Excel report + every uploaded document) from the saved data.
            const payload = buildPayload();
            const encoded = encodeURIComponent(JSON.stringify(payload));
            if (encoded.length > 60000) {
                alert('Too much data to save before building the package. Press Save All to Cloud first.');
                return;
            }
            persistentStoreData[payload.spk] = payload;
            try { localStorage.setItem('oz_data_' + payload.spk, JSON.stringify(payload)); } catch(e) {}
            navigateParent(ctxQS('&action=save_all&then=package&payload=' + encoded));
        }

        function openAddUserModal() {
            document.getElementById('edit-form-title').textContent = "Add New User Credential";
            document.getElementById('edit-orig-username').value = "";
            document.getElementById('edit-username').value = "";
            document.getElementById('edit-name').value = "";
            document.getElementById('edit-role').value = "store";
            document.getElementById('edit-password').value = "";
            document.getElementById('user-edit-form-box').classList.remove('hidden');
        }

        function editUser(username, name, role, password) {
            document.getElementById('edit-form-title').textContent = `Edit User: ${username}`;
            document.getElementById('edit-orig-username').value = username;
            document.getElementById('edit-username').value = username;
            document.getElementById('edit-name').value = name;
            document.getElementById('edit-role').value = role;
            document.getElementById('edit-password').value = password;
            document.getElementById('user-edit-form-box').classList.remove('hidden');
        }

        function submitUserEdit(event) {
            event.preventDefault();
            const orig = document.getElementById('edit-orig-username').value;
            const user = document.getElementById('edit-username').value;
            const name = document.getElementById('edit-name').value;
            const role = document.getElementById('edit-role').value;
            const pass = document.getElementById('edit-password').value;

            navigateParent(`?auth_user=${encodeURIComponent(currentUsername)}&action=update_user&orig_username=${encodeURIComponent(orig)}&username=${encodeURIComponent(user)}&name=${encodeURIComponent(name)}&role=${encodeURIComponent(role)}&password=${encodeURIComponent(pass)}`);
        }

        function deleteUser(username) {
            if (confirm(`Are you sure you want to delete user "${username}"?`)) {
                navigateParent(`?auth_user=${encodeURIComponent(currentUsername)}&action=delete_user&username=${encodeURIComponent(username)}`);
            }
        }

        function updateFinancialSummary() {
            try {
                const spk = getSPK();
                const inv = storeInventoryData[spk];
                let totalSales = 0;
                let totalCogs = 0;

                if (inv && inv.rows && inv.headers) {
                    const soldHeader = inv.headers.find(h => h && h.toLowerCase().includes('sold'));
                    const costHeader = inv.headers.find(h => h && (h.toLowerCase().includes('cost') || h.toLowerCase().includes('avg cost')));

                    inv.rows.forEach(row => {
                        let soldVal = 0;
                        if (soldHeader && row[soldHeader] !== undefined) {
                            soldVal = parseFloat(String(row[soldHeader]).replace(/[^0-9.-]/g, '')) || 0;
                        } else {
                            const keys = Object.keys(row).filter(k => k !== 'Category');
                            if (keys.length > 0) soldVal = parseFloat(String(row[keys[0]]).replace(/[^0-9.-]/g, '')) || 0;
                        }
                        totalSales += soldVal;

                        if (costHeader && row[costHeader] !== undefined) {
                            const costVal = parseFloat(String(row[costHeader]).replace(/[^0-9.-]/g, '')) || 0;
                            totalCogs += costVal;
                        }
                    });
                }

                if (totalCogs === 0 && totalSales > 0) {
                    totalCogs = totalSales * 0.80;
                }

                const parseNum = (id) => {
                    const el = document.getElementById(id);
                    if (!el) return 0;
                    return parseFloat(String(el.value).replace(/[^0-9.-]/g, '')) || 0;
                };

                const expRent = parseNum('exp-rent');
                const expSal = parseNum('exp-salaries');
                const expUtil = parseNum('exp-utilities');
                const expAdv = parseNum('exp-advertising');
                const expMkt = parseNum('exp-marketing');
                const expData = parseNum('exp-datadeals');
                const expIns = parseNum('exp-insurance');
                const expGar = parseNum('exp-garbage');
                const expSec = parseNum('exp-security');
                const expLic = parseNum('exp-licenses');

                const totalExp = expRent + expSal + expUtil + expAdv + expMkt + expData + expIns + expGar + expSec + expLic;

                const grossProfit = totalSales - totalCogs;
                const grossMargin = totalSales > 0 ? (grossProfit / totalSales) * 100 : 0;

                const netOperatingIncome = grossProfit - totalExp;
                const operatingMargin = totalSales > 0 ? (netOperatingIncome / totalSales) * 100 : 0;

                const formatMoney = (n) => '$' + n.toLocaleString(undefined, { minimumFractionDigits: 2, maximumFractionDigits: 2 });
                const formatPct = (n) => n.toFixed(1) + '%';

                const salesEl = document.getElementById('fin-total-sales');
                const cogsEl = document.getElementById('fin-cogs');
                const grossMEl = document.getElementById('fin-gross-margin');
                const netEl = document.getElementById('fin-net');
                const netMEl = document.getElementById('fin-net-margin');
                const totExpEl = document.getElementById('fin-total-expenses');

                if (salesEl) salesEl.textContent = formatMoney(totalSales);
                if (cogsEl) cogsEl.textContent = formatMoney(totalCogs);
                if (grossMEl) grossMEl.textContent = 'Gross Margin: ' + formatPct(grossMargin);
                if (netEl) {
                    netEl.textContent = formatMoney(netOperatingIncome);
                    netEl.className = netOperatingIncome >= 0 ? 'text-lg font-bold text-white' : 'text-lg font-bold text-red-300';
                }
                if (netMEl) {
                    netMEl.textContent = formatPct(operatingMargin);
                    netMEl.className = operatingMargin >= 0 ? 'text-lg font-bold text-amber-300' : 'text-lg font-bold text-red-300';
                }
                if (totExpEl) totExpEl.textContent = formatMoney(totalExp);
            } catch(e) {
                console.error("updateFinancialSummary error:", e);
            }
        }

        initApp();
    </script>
</body>
</html>"""


# ============================================================================
# STREAMLIT APP  (original HTML interface + real file storage + report engine)
# ============================================================================
import hashlib
import hmac
import html as htmllib
import inspect
import json
import logging
import os
import secrets as pysecrets
import sqlite3
import tempfile
import uuid
from contextlib import closing
from datetime import date
from pathlib import Path

import streamlit as st
import streamlit.components.v1 as components

# Suppress deprecation warning for components.html
logging.getLogger("streamlit.runtime.scriptrunner_utils.script_run_context").setLevel(logging.ERROR)
log = logging.getLogger("oz_store")

# --- Page config ---
st.set_page_config(
    page_title="The OZ Store - Persistent Accountant Package Manager",
    page_icon="🌿",
    layout="wide",
)

hide_streamlit_style = """
    <style>
    #MainMenu {visibility: hidden;}
    footer {visibility: hidden;}
    header {visibility: hidden;}
    .block-container { padding: 0rem; }
    [data-testid="stExpander"], [data-testid="stAlert"] { margin: 0.75rem 1.25rem 0 1.25rem; }
    .stButton>button, .stDownloadButton>button { background-color: #047857; color: white; border-radius: 12px; font-weight: bold; border: none; }
    .stButton>button:hover, .stDownloadButton>button:hover { background-color: #065f46; color: white; }
    </style>
"""
st.markdown(hide_streamlit_style, unsafe_allow_html=True)


def _stretch():
    """Full-width widgets: newer Streamlit uses width='stretch', older use_container_width."""
    try:
        if "width" in inspect.signature(st.button).parameters:
            return {"width": "stretch"}
    except Exception:
        pass
    return {"use_container_width": True}


STRETCH = _stretch()

# ----------------------------------------------------------------------------
# Config
# ----------------------------------------------------------------------------
DROPBOX_CREDENTIALS_FILE = "dropbox_credentials.json"
DROPBOX_FOLDER_PATH = "/OZ_Store_Uploads"
MASTER_FILE = "master_store_state.json"
USERS_FILE = "users.json"
DB_FILE = "oz_store.db"
# NOTE: not called "uploads": a leftover *file* with that name crashed the old app
LOCAL_FILES_DIR = Path(tempfile.gettempdir()) / "oz_store_files"
MAX_UPLOAD_MB = 25

STORES = [f"Oz {i}" for i in range(1, 8)]
MONTHS = ["January", "February", "March", "April", "May", "June", "July",
          "August", "September", "October", "November", "December"]
YEARS = ["2025", "2026", "2027", "2028"]
ROLES = ("store", "accountant", "admin")

UPLOAD_CATS = [
    ("cooh", "Cash On Hand - End of Day"),
    ("brinks", "Brinks Book"),
    ("green1", "Greenline: Inventory Summary"),
    ("green2", "Greenline: Financials"),
    ("green3", "Greenline: Payment Types"),
    ("mon1", "Moneris: Counter Sales"),
    ("mon2", "Moneris: Online Sales"),
    ("ocs1", "OCS: Account Statement"),
    ("ocs2", "OCS: Invoices"),
    ("chargebacks", "Moneris Chargebacks"),
    ("gas1", "Gas Cards: Details"),
    ("gas2", "Gas Cards: Summary"),
]
DOC_LABELS = dict(UPLOAD_CATS)

DEFAULT_USERS = [
    ("admin", "Master Admin", "admin", "admin2026"),
    ("leanne", "Leanne (Accountant)", "accountant", "leanne2026"),
    ("oz1", "Oz 1", "store", "oz1pass2026"),
    ("oz2", "Oz 2", "store", "oz2pass2026"),
    ("oz3", "Oz 3", "store", "oz3pass2026"),
    ("oz4", "Oz 4", "store", "oz4pass2026"),
    ("oz5", "Oz 5", "store", "oz5pass2026"),
    ("oz6", "Oz 6", "store", "oz6pass2026"),
    ("oz7", "Oz 7", "store", "oz7pass2026"),
]


def secret(name, default=None):
    try:
        return st.secrets.get(name, default)
    except Exception:
        return default


# ----------------------------------------------------------------------------
# Dropbox
# ----------------------------------------------------------------------------
def get_dropbox_client():
    """Refresh-token login (never expires) if configured, else the access token."""
    try:
        import dropbox
    except ImportError:
        return None
    try:
        key, sec, rt = secret("DROPBOX_APP_KEY"), secret("DROPBOX_APP_SECRET"), secret("DROPBOX_REFRESH_TOKEN")
        if key and sec and rt:
            return dropbox.Dropbox(oauth2_refresh_token=rt, app_key=key, app_secret=sec)
        token = secret("DROPBOX_ACCESS_TOKEN")
        if token:
            return dropbox.Dropbox(token)
        if os.path.exists(DROPBOX_CREDENTIALS_FILE):
            with open(DROPBOX_CREDENTIALS_FILE, "r") as f:
                token = json.load(f).get("access_token")
            if token:
                return dropbox.Dropbox(token)
    except Exception:
        log.exception("Dropbox client could not be created")
    return None


def _dbx_put(dbx, path: str, data: bytes):
    import dropbox
    dbx.files_upload(data, path, mode=dropbox.files.WriteMode.overwrite)


def _dbx_get(dbx, path: str) -> bytes:
    _, res = dbx.files_download(path)
    return res.content


def _dbx_missing(exc) -> bool:
    try:
        return exc.error.is_path() and exc.error.get_path().is_not_found()
    except Exception:
        return False


# ----------------------------------------------------------------------------
# Database (same tables and columns as before, so existing data keeps working)
# ----------------------------------------------------------------------------
def _connect():
    return closing(sqlite3.connect(DB_FILE))


def _loads(s, default):
    try:
        v = json.loads(s) if s else default
        return default if v is None else v
    except Exception:
        return default


def load_users_from_dropbox():
    dbx = get_dropbox_client()
    if not dbx:
        return None
    try:
        users = json.loads(_dbx_get(dbx, f"{DROPBOX_FOLDER_PATH}/{USERS_FILE}"))
        ok = [tuple(u) for u in users if isinstance(u, list) and len(u) == 4]
        return ok or None
    except Exception:
        return None


def sync_users_to_dropbox():
    dbx = get_dropbox_client()
    if not dbx:
        return
    try:
        with _connect() as conn:
            rows = conn.execute("SELECT username, name, role, password FROM users").fetchall()
        _dbx_put(dbx, f"{DROPBOX_FOLDER_PATH}/{USERS_FILE}", json.dumps([list(r) for r in rows]).encode())
    except Exception:
        log.exception("Could not back up users to Dropbox")


def init_db():
    try:
        LOCAL_FILES_DIR.mkdir(parents=True, exist_ok=True)
    except OSError:
        log.exception("Could not create local file folder")
    with _connect() as conn:
        cur = conn.cursor()
        cur.execute("CREATE TABLE IF NOT EXISTS users (username TEXT PRIMARY KEY, name TEXT, role TEXT, password TEXT)")
        cur.execute(
            "CREATE TABLE IF NOT EXISTS store_data (store_period_key TEXT PRIMARY KEY, store TEXT, period TEXT, "
            "cooh_text TEXT, cooh_date TEXT, chargebacks_text TEXT, checklist_json TEXT, inventory_json TEXT, "
            "inventory_filename TEXT)"
        )
        cur.execute(
            "CREATE TABLE IF NOT EXISTS store_files (id INTEGER PRIMARY KEY AUTOINCREMENT, store_period_key TEXT, "
            "category TEXT, filename TEXT, filepath TEXT)"
        )
        # Add any column an older version of the app didn't have
        existing = {c[1] for c in cur.execute("PRAGMA table_info(store_data)")}
        for col in ["store", "period", "cooh_text", "cooh_date", "chargebacks_text",
                    "checklist_json", "inventory_json", "inventory_filename"]:
            if col not in existing:
                cur.execute(f"ALTER TABLE store_data ADD COLUMN {col} TEXT")
        if cur.execute("SELECT COUNT(*) FROM users").fetchone()[0] == 0:
            cur.executemany("INSERT OR IGNORE INTO users VALUES (?, ?, ?, ?)",
                            load_users_from_dropbox() or DEFAULT_USERS)
        conn.commit()


def get_user(username):
    with _connect() as conn:
        return conn.execute("SELECT name, role, password FROM users WHERE username = ?",
                            ((username or "").lower().strip(),)).fetchone()


def normalize_inventory(inv):
    """The page expects {headers, rows}. Accept the other shapes older versions saved."""
    if isinstance(inv, dict) and isinstance(inv.get("rows"), list):
        rows = [r for r in inv["rows"] if isinstance(r, dict)]
        headers = inv.get("headers") or (list(rows[0].keys()) if rows else [])
        return {"headers": headers, "rows": rows} if rows else None
    if isinstance(inv, list) and inv and isinstance(inv[0], dict):
        return {"headers": list(inv[0].keys()), "rows": inv}
    return None


def normalize_remote(d: dict) -> dict:
    """Bring any saved format (including the short Streamlit-widget versions) to the page format."""
    cj = d.get("checklist_json")
    if not isinstance(cj, dict):
        notes = dict(d.get("notes") or {})
        cj = {"checkboxes": d.get("checklist") or {}, "notes": notes,
              "dates": {}, "expenses": d.get("expenses") or {}, "files": {}}
        cj["notes"].setdefault("cooh", d.get("cooh_text", ""))
        cj["notes"].setdefault("chargebacks", d.get("chargebacks_text", ""))
    cj.setdefault("files", {})
    if not isinstance(cj["files"], dict):
        cj["files"] = {}
    return {
        "cooh_text": d.get("cooh_text") or "",
        "cooh_date": d.get("cooh_date") or "",
        "chargebacks_text": d.get("chargebacks_text") or "Nothing to report.",
        "checklist_json": cj,
        "inventory_data": normalize_inventory(d.get("inventory_data") or d.get("inventory_json")),
        "inventory_filename": d.get("inventory_filename") or "",
        "_updated": d.get("_updated") or cj.get("_updated") or 0,
    }


_REC_COLS = ("store_period_key, cooh_text, cooh_date, chargebacks_text, checklist_json, "
             "inventory_json, inventory_filename")


def _row_to_rec(r) -> dict:
    cj = _loads(r[4], {})
    return normalize_remote({
        "cooh_text": r[1], "cooh_date": r[2], "chargebacks_text": r[3], "checklist_json": cj,
        "inventory_data": _loads(r[5], None), "inventory_filename": r[6],
        "_updated": cj.get("_updated", 0) if isinstance(cj, dict) else 0,
    })


def db_all() -> dict:
    with _connect() as conn:
        rows = conn.execute(f"SELECT {_REC_COLS} FROM store_data").fetchall()
    return {r[0]: _row_to_rec(r) for r in rows}


def get_record(spk: str):
    with _connect() as conn:
        r = conn.execute(f"SELECT {_REC_COLS} FROM store_data WHERE store_period_key = ?", (spk,)).fetchone()
    return _row_to_rec(r) if r else None


def blank_record() -> dict:
    return normalize_remote({"checklist_json": {"checkboxes": {}, "notes": {}, "dates": {}, "expenses": {}, "files": {}}})


def save_record(spk: str, rec: dict):
    store, _, rest = spk.partition("_")
    cj = dict(rec.get("checklist_json") or {})
    cj["_updated"] = rec.get("_updated") or cj.get("_updated") or 0
    inv = normalize_inventory(rec.get("inventory_data"))
    with _connect() as conn:
        conn.execute(
            "INSERT OR REPLACE INTO store_data (store_period_key, store, period, cooh_text, cooh_date, "
            "chargebacks_text, checklist_json, inventory_json, inventory_filename) VALUES (?,?,?,?,?,?,?,?,?)",
            (spk, store, rest.replace("_", " "), rec.get("cooh_text", ""), rec.get("cooh_date", ""),
             rec.get("chargebacks_text", "Nothing to report."), json.dumps(cj),
             json.dumps(inv) if inv else None, rec.get("inventory_filename", "")),
        )
        conn.commit()


def sync_all_to_dropbox():
    """Write the master JSON that the phone reads. Returns (ok, error text)."""
    dbx = get_dropbox_client()
    if not dbx:
        return False, "Dropbox isn't configured in Streamlit secrets, so data is only on this server."
    try:
        _dbx_put(dbx, f"{DROPBOX_FOLDER_PATH}/{MASTER_FILE}",
                 json.dumps(db_all(), indent=2, default=str).encode("utf-8"))
        return True, ""
    except Exception as e:
        log.exception("Dropbox sync failed")
        return False, str(e)


@st.cache_data(ttl=120, show_spinner=False)
def restore_missing_from_dropbox():
    """After a reboot the server database is empty: pull anything missing back from Dropbox."""
    dbx = get_dropbox_client()
    if not dbx:
        return ""
    try:
        remote = json.loads(_dbx_get(dbx, f"{DROPBOX_FOLDER_PATH}/{MASTER_FILE}"))
    except Exception:
        return ""
    if not isinstance(remote, dict):
        return ""
    with _connect() as conn:
        have = {r[0] for r in conn.execute("SELECT store_period_key FROM store_data")}
    for spk, d in remote.items():
        if spk not in have and isinstance(d, dict):
            try:
                save_record(spk, normalize_remote(d))
            except Exception:
                log.exception("Could not restore %s", spk)
    return "ok"


# ----------------------------------------------------------------------------
# Uploaded documents: stored for real (Dropbox, or this server as a fallback)
# ----------------------------------------------------------------------------
def safe_name(n: str) -> str:
    return (re.sub(r"[^\w.\- ]", "_", n or "").strip() or "file")[:120]


def store_file(spk: str, cat: str, name: str, data: bytes):
    """Returns (meta, warning). meta is what we keep in the checklist record."""
    fid = uuid.uuid4().hex[:10]
    fname = f"{fid}_{safe_name(name)}"
    meta = {"id": fid, "filename": name, "size": len(data), "uploaded": datetime.now().isoformat(timespec="seconds")}
    warn = None
    dbx = get_dropbox_client()
    if dbx:
        path = f"{DROPBOX_FOLDER_PATH}/files/{spk}/{cat}/{fname}"
        try:
            _dbx_put(dbx, path, data)
            return {**meta, "backend": "dropbox", "path": path}, None
        except Exception as e:
            log.exception("Dropbox upload failed")
            warn = f"Dropbox upload failed ({e}); kept on the server only."
    p = LOCAL_FILES_DIR / spk / cat / fname
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_bytes(data)
    return {**meta, "backend": "local", "path": str(p)}, warn


def read_file_bytes(meta: dict):
    try:
        if meta.get("backend") == "dropbox":
            dbx = get_dropbox_client()
            return _dbx_get(dbx, meta["path"]) if dbx else None
        p = Path(meta.get("path", ""))
        return p.read_bytes() if p.is_file() else None
    except Exception:
        log.exception("Could not read %s", meta.get("filename"))
        return None


def delete_file_bytes(meta: dict):
    try:
        if meta.get("backend") == "dropbox":
            dbx = get_dropbox_client()
            if dbx:
                dbx.files_delete_v2(meta["path"])
        else:
            Path(meta.get("path", "")).unlink(missing_ok=True)
    except Exception:
        log.exception("Could not delete %s", meta.get("filename"))


def add_files_to_record(spk: str, cat: str, metas: list):
    rec = get_record(spk) or blank_record()
    files = rec["checklist_json"].setdefault("files", {})
    files.setdefault(cat, []).extend(metas)
    save_record(spk, rec)


# ----------------------------------------------------------------------------
# Report engine glue: saved record -> engine payload -> Excel + ZIP
# ----------------------------------------------------------------------------
def rec_to_engine(rec: dict) -> dict:
    cj = rec.get("checklist_json") or {}
    inv = normalize_inventory(rec.get("inventory_data"))
    recs = None
    if inv:
        headers = inv["headers"]
        recs = [{h: row.get(h) for h in headers} for row in inv["rows"]]
    files = {}
    for cat, lst in (cj.get("files") or {}).items():
        if lst:
            files[cat] = [{"id": f.get("id"), "name": f.get("filename", ""), "size": f.get("size", 0), "meta": f}
                          for f in lst]
    return {
        "checklist": cj.get("checkboxes") or {}, "notes": cj.get("notes") or {},
        "cooh_text": rec.get("cooh_text", ""), "cooh_date": rec.get("cooh_date", ""),
        "chargebacks_text": rec.get("chargebacks_text", ""), "expenses": cj.get("expenses") or {},
        "inventory_data": recs, "inventory_filename": rec.get("inventory_filename", ""), "files": files,
    }


def gather_documents(payload: dict):
    docs, unreadable = [], []
    for cat, lst in (payload.get("files") or {}).items():
        for f in lst:
            data = read_file_bytes(f["meta"])
            label = DOC_LABELS.get(cat, cat)
            if data is None:
                unreadable.append((label, f["name"]))
            else:
                docs.append({"key": f"{cat}:{f['id']}", "label": label, "name": f["name"], "data": data})
    return docs, unreadable


@st.cache_data(show_spinner=False, max_entries=64)
def parse_table(data: bytes, name: str):
    try:
        df = read_table(name, data)
        return _coerce_numeric(df.copy()) if df is not None else None
    except Exception:
        log.exception("Could not parse %s", name)
        return None


def package_entries(store: str, period: str, payload: dict):
    docs, unreadable = gather_documents(payload)
    tables, doc_log = {}, []
    for d in docs:
        df = parse_table(d["data"], d["name"])
        kb = round(len(d["data"]) / 1024, 1)
        if df is not None and not df.empty:
            tables[d["key"]] = (d["label"], df)
            doc_log.append([d["label"], d["name"], kb, f"Own sheet ({len(df)} rows)"])
        else:
            doc_log.append([d["label"], d["name"], kb, "Attached only (PDF, image or empty)"])
    for label, name in unreadable:
        doc_log.append([label, name, "", "NOT FOUND in storage - please upload again"])
    xlsx = build_store_workbook(store, period, payload, tables, doc_log)
    inv_rows = payload.get("inventory_data")
    base = safe_name(f"{store} - {period}")
    entries, seen = [(f"{base}/{base} - Report.xlsx", xlsx)], set()
    for d in docs:
        arc = f"{base}/Documents/{safe_name(d['label'])} - {safe_name(d['name'])}"
        n = 2
        while arc in seen:
            arc = f"{base}/Documents/{n} - {safe_name(d['label'])} - {safe_name(d['name'])}"
            n += 1
        seen.add(arc)
        entries.append((arc, d["data"]))
    if inv_rows:
        entries.append((f"{base}/Documents/Greenline Inventory (formatted).csv",
                        pd.DataFrame(inv_rows).to_csv(index=False).encode("utf-8-sig")))
    return xlsx, entries, unreadable


def build_store_package(store: str, month: str, year: str):
    spk = f"{store}_{month}_{year}"
    rec = get_record(spk)
    if not rec:
        return None
    xlsx, entries, unreadable = package_entries(store, f"{month} {year}", rec_to_engine(rec))
    return {"spk": spk, "xlsx": xlsx, "zip": build_zip(entries), "n_docs": len(entries) - 1,
            "unreadable": unreadable, "at": datetime.now().strftime("%H:%M:%S")}


def build_all_package(month: str, year: str):
    period, payloads, entries = f"{month} {year}", {}, []
    for s in STORES:
        rec = get_record(f"{s}_{month}_{year}")
        if rec:
            payloads[s] = rec_to_engine(rec)
            entries += package_entries(s, period, payloads[s])[1]
    if not payloads:
        return None
    summary = build_consolidated_workbook(period, payloads)
    entries.insert(0, (f"All Stores - {period} - Summary.xlsx", summary))
    return {"key": f"{month}_{year}", "xlsx": summary, "zip": build_zip(entries), "stores": len(payloads)}


# ----------------------------------------------------------------------------
# Login tokens (the page keeps you signed in through the URL; this makes it unforgeable)
# ----------------------------------------------------------------------------
@st.cache_resource
def _auth_key() -> bytes:
    return str(secret("AUTH_SECRET") or pysecrets.token_hex(32)).encode()


def make_token(username: str, password: str) -> str:
    sig = hmac.new(_auth_key(), f"{username}:{password}".encode(), hashlib.sha256).hexdigest()[:32]
    return f"{username}~{sig}"


def user_from_token(token: str):
    """-> (username, name, role) or None"""
    if not token or "~" not in token:
        return None
    username, _, sig = token.partition("~")
    rec = get_user(username)
    if not rec:
        return None
    if hmac.compare_digest(make_token(username.lower().strip(), rec[2]), f"{username.lower().strip()}~{sig}"):
        return username.lower().strip(), rec[0], rec[1]
    return None


# ----------------------------------------------------------------------------
# Boot
# ----------------------------------------------------------------------------
init_db()
restore_missing_from_dropbox()

for _k, _v in {"logged_in": False, "username": None, "name": None, "role": None}.items():
    st.session_state.setdefault(_k, _v)

query_params = st.query_params


def _log_in(username, name, role):
    st.session_state.logged_in = True
    st.session_state.username = username
    st.session_state.name = name
    st.session_state.role = role


def _log_out():
    for k in ("logged_in", "username", "name", "role", "ctx", "pkg", "all_pkg"):
        st.session_state.pop(k, None)


if not st.session_state.logged_in and "auth_user" in query_params:
    _u = user_from_token(query_params.get("auth_user", ""))
    if _u:
        _log_in(*_u)

# ----------------------------------------------------------------------------
# Login screen
# ----------------------------------------------------------------------------
if not st.session_state.logged_in:
    col1, col2, col3 = st.columns([1, 1.2, 1])
    with col2:
        st.markdown("<br><br>", unsafe_allow_html=True)
        st.markdown(
            """
            <div style="background-color: #065f46; padding: 30px; border-radius: 16px; color: white; text-align: center; box-shadow: 0 10px 25px rgba(0,0,0,0.1);">
                <div style="font-size: 40px; margin-bottom: 10px;">🌿</div>
                <h2 style="margin: 0; font-size: 22px;">The OZ Store Portal</h2>
                <p style="font-size: 12px; color: #a7f3d0; margin-top: 5px;">Secure Persistent Database Manager</p>
            </div>
        """,
            unsafe_allow_html=True,
        )

        with st.form("login_form"):
            username_input = st.text_input("Username")
            password_input = st.text_input("Password", type="password")
            submit = st.form_submit_button("Sign In", **STRETCH)

            if submit:
                user_record = get_user(username_input)
                if user_record and hmac.compare_digest(str(user_record[2]).encode(), password_input.encode()):
                    clean_u = username_input.lower().strip()
                    _log_in(clean_u, user_record[0], user_record[1])
                    st.query_params.clear()
                    st.query_params["auth_user"] = make_token(clean_u, user_record[2])
                    st.rerun()
                else:
                    st.error("Invalid username or password.")
    st.stop()

role = st.session_state.role
is_staff = role in ("admin", "accountant")


# ----------------------------------------------------------------------------
# Which store / month / year is on screen (set by the page's tabs and period menu)
# ----------------------------------------------------------------------------
def allowed_stores():
    return [st.session_state.name] if role == "store" else STORES


def init_ctx():
    today = date.today()
    ctx = st.session_state.get("ctx") or {
        "store": None, "month": MONTHS[today.month - 1],
        "year": str(today.year) if str(today.year) in YEARS else YEARS[1],
    }
    allowed = allowed_stores()
    s, m, y = query_params.get("store"), query_params.get("month"), query_params.get("year")
    if s in allowed:
        ctx["store"] = s
    if m in MONTHS:
        ctx["month"] = m
    if y in YEARS:
        ctx["year"] = y
    if ctx["store"] not in allowed:
        ctx["store"] = allowed[0]
    st.session_state.ctx = ctx
    return ctx


ctx = init_ctx()
spk = f"{ctx['store']}_{ctx['month']}_{ctx['year']}"


def current_token():
    rec = get_user(st.session_state.username)
    return make_token(st.session_state.username, rec[2]) if rec else None


def reset_url_and_rerun():
    """Drop the action from the URL but keep you signed in on the same store/period."""
    tok = current_token()
    if not tok:
        _log_out()
        st.query_params.clear()
        st.rerun()
    c = st.session_state.ctx
    params = {"auth_user": tok, "store": c["store"], "month": c["month"], "year": c["year"]}
    st.query_params.clear()
    if hasattr(st.query_params, "from_dict"):
        st.query_params.from_dict(params)
    else:
        for k, v in params.items():
            st.query_params[k] = v
    st.rerun()


def flash(kind: str, text: str):
    st.session_state.flash = (kind, text)


# ----------------------------------------------------------------------------
# Actions sent by the page (these arrive as URL parameters)
# ----------------------------------------------------------------------------
def act_save_all():
    raw = query_params.get("payload", "")
    if not raw or len(raw) > 400_000:
        return flash("error", "Nothing was saved: the data was empty or too large.")
    try:
        payload = json.loads(raw)
    except ValueError:
        try:
            payload = json.loads(urllib.parse.unquote(raw))
        except ValueError:
            return flash("error", "Nothing was saved: the data could not be read.")
    store = payload.get("store", "")
    month, _, year = str(payload.get("period", "")).partition(" ")
    if store not in allowed_stores() or month not in MONTHS or year not in YEARS:
        return flash("error", "Nothing was saved: that store or period isn't allowed for your account.")
    target = f"{store}_{month}_{year}"
    old = get_record(target) or blank_record()
    cj = payload.get("checklist_json") if isinstance(payload.get("checklist_json"), dict) else {}
    cj["files"] = old["checklist_json"].get("files", {})  # files are managed on the server only
    rec = normalize_remote({
        "cooh_text": payload.get("cooh_text", ""), "cooh_date": payload.get("cooh_date", ""),
        "chargebacks_text": payload.get("chargebacks_text", "Nothing to report."), "checklist_json": cj,
        "inventory_data": payload.get("inventory_data") or payload.get("inventory_json"), "inventory_filename": payload.get("inventory_filename", ""),
        "_updated": payload.get("_updated") or 0,
    })
    save_record(target, rec)
    ok, err = sync_all_to_dropbox()
    if query_params.get("then") == "package":
        st.session_state.panel_open = True
        st.session_state.auto_pkg = True
    if ok:
        flash("success", f"Saved {store} - {month} {year} and synced to Dropbox.")
    else:
        flash("warning", f"Saved on the server, but Dropbox sync failed: {err}")


def act_update_user():
    if role != "admin":
        return flash("error", "Only the admin can change users.")
    orig = query_params.get("orig_username", "").lower().strip()
    user = query_params.get("username", "").lower().strip()
    name = query_params.get("name", "").strip()
    new_role = query_params.get("role", "")
    password = query_params.get("password", "")
    if not re.fullmatch(r"[a-z0-9_.-]{2,30}", user) or not name or len(name) > 60 or not password or len(password) > 100:
        return flash("error", "User not saved: check the username (letters/numbers), display name and password.")
    if new_role not in ROLES:
        return flash("error", "User not saved: unknown role.")
    if new_role == "store" and name not in STORES:
        return flash("error", "User not saved: a store login's display name must be exactly one of " + ", ".join(STORES) + ".")
    with _connect() as conn:
        if orig and orig != user:
            conn.execute("DELETE FROM users WHERE username = ?", (orig,))
        conn.execute("INSERT OR REPLACE INTO users VALUES (?, ?, ?, ?)", (user, name, new_role, password))
        conn.commit()
    if orig and orig == st.session_state.username and orig != user:
        st.session_state.username = user
    sync_users_to_dropbox()
    flash("success", f"User {user} saved.")


def act_delete_user():
    if role != "admin":
        return flash("error", "Only the admin can change users.")
    user = query_params.get("username", "").lower().strip()
    if user == st.session_state.username:
        return flash("error", "You can't delete the account you're signed in with.")
    with _connect() as conn:
        conn.execute("DELETE FROM users WHERE username = ?", (user,))
        conn.commit()
    sync_users_to_dropbox()
    flash("success", f"User {user} deleted.")


def act_delete_file():
    cat, fid = query_params.get("cat", ""), query_params.get("file_id", "")
    rec = get_record(spk)
    files = ((rec or {}).get("checklist_json") or {}).get("files", {})
    keep, gone = [], []
    for f in files.get(cat, []):
        (gone if f.get("id") == fid else keep).append(f)
    if not gone:
        return flash("warning", "That file was already removed.")
    for f in gone:
        delete_file_bytes(f)
    files[cat] = keep
    save_record(spk, rec)
    sync_all_to_dropbox()
    flash("success", f"Removed {gone[0].get('filename', 'file')}.")


def act_clear_files():
    cat = query_params.get("cat", "")
    rec = get_record(spk)
    files = ((rec or {}).get("checklist_json") or {}).get("files", {})
    for f in files.get(cat, []):
        delete_file_bytes(f)
    if files.get(cat):
        files[cat] = []
        save_record(spk, rec)
        sync_all_to_dropbox()
    flash("success", f"Cleared all files in {DOC_LABELS.get(cat, cat)}.")


def handle_action(action: str):
    if action == "sign_out":
        _log_out()
        st.query_params.clear()
        st.rerun()
    elif action == "save_all":
        act_save_all()
    elif action == "update_user":
        act_update_user()
    elif action == "delete_user":
        act_delete_user()
    elif action == "delete_file":
        act_delete_file()
    elif action == "clear_files":
        act_clear_files()
    elif action == "upload_panel":
        cat = query_params.get("cat", "")
        if cat in DOC_LABELS:
            st.session_state.upl_cat_pref = cat
        st.session_state.panel_open = True
    elif action == "package":
        st.session_state.panel_open = True
        st.session_state.auto_pkg = True
    reset_url_and_rerun()


if "action" in query_params:
    handle_action(query_params.get("action", ""))

# ----------------------------------------------------------------------------
# Documents & reports panel (real uploads + the report generator)
# ----------------------------------------------------------------------------
def render_panel():
    store, month, year = ctx["store"], ctx["month"], ctx["year"]
    pref = st.session_state.pop("upl_cat_pref", None)
    if pref in DOC_LABELS:
        st.session_state["upl_cat"] = pref
    st.session_state.setdefault("upl_ver", 0)
    open_panel = bool(st.session_state.pop("panel_open", False))

    with st.expander(f"📎 Documents & Reports — {store} · {month} {year}", expanded=open_panel):
        st.caption("Files attached here are stored for real (Dropbox) and included in the accountant package. "
                   "Use the store tabs and period menu below to change store or month.")
        c1, c2 = st.columns([1, 2])
        with c1:
            cat = st.selectbox("Document category", [k for k, _ in UPLOAD_CATS],
                               format_func=lambda k: DOC_LABELS[k], key="upl_cat")
        with c2:
            files = st.file_uploader("Choose one or more files (CSV, Excel, PDF, photos)",
                                     accept_multiple_files=True, key=f"upl_files_{st.session_state.upl_ver}")
        if st.button("⬆️ Upload selected files", disabled=not files, **STRETCH):
            ok_n, problems = 0, []
            metas = []
            for f in files:
                if f.size > MAX_UPLOAD_MB * 1024 * 1024:
                    problems.append(f"{f.name} is over {MAX_UPLOAD_MB} MB.")
                    continue
                meta, warn = store_file(spk, cat, f.name, f.getvalue())
                metas.append(meta)
                ok_n += 1
                if warn:
                    problems.append(warn)
            if metas:
                add_files_to_record(spk, cat, metas)
                synced, err = sync_all_to_dropbox()
                if not synced:
                    problems.append(f"Dropbox sync failed: {err}")
            st.session_state.upl_ver += 1
            st.session_state.panel_open = True
            if problems:
                flash("warning", f"{ok_n} file(s) saved to {DOC_LABELS[cat]}. " + " ".join(dict.fromkeys(problems)))
            else:
                flash("success", f"{ok_n} file(s) saved to {DOC_LABELS[cat]}. Tick its box in the checklist and press Save All.")
            st.rerun()

        st.markdown("**Accountant package**")
        st.caption("Builds from the last saved data (press Save All to Cloud first): an Excel report with the checklist, "
                   "cash, expenses, financial summary, formatted inventory and one sheet for every CSV/Excel document, "
                   "plus a ZIP containing the report and all the original files.")
        auto = bool(st.session_state.pop("auto_pkg", False))
        if st.button("⚙️ Build package for this store & month", key="build_pkg", **STRETCH) or auto:
            with st.spinner("Reading documents and building the report…"):
                pkg = build_store_package(store, month, year)
            if pkg is None:
                st.info("Nothing saved yet for this store and month. Press Save All to Cloud first.")
            else:
                st.session_state.pkg = pkg
        pkg = st.session_state.get("pkg")
        if pkg and pkg["spk"] == spk:
            base = safe_name(f"{store} - {month} {year}")
            st.success(f"Package ready: {pkg['n_docs']} document(s) included (built {pkg['at']}).")
            for label, name in pkg["unreadable"]:
                st.warning(f"Couldn't read {name} ({label}) from storage; please upload it again.")
            d1, d2 = st.columns(2)
            d1.download_button("⬇️ Excel report", pkg["xlsx"], file_name=f"{base} - Report.xlsx",
                               mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                               key="dl_xlsx", **STRETCH)
            d2.download_button("⬇️ Full package (ZIP)", pkg["zip"], file_name=f"{base} - Package.zip",
                               mime="application/zip", key="dl_zip", **STRETCH)

        if is_staff:
            st.divider()
            st.markdown(f"**All stores — {month} {year}**")
            if st.button("📚 Build consolidated package (all stores)", key="build_all", **STRETCH):
                with st.spinner("Building every store…"):
                    allp = build_all_package(month, year)
                if allp is None:
                    st.info("No store has saved data for this month yet.")
                else:
                    st.session_state.all_pkg = allp
            allp = st.session_state.get("all_pkg")
            if allp and allp["key"] == f"{month}_{year}":
                st.success(f"Consolidated package ready ({allp['stores']} store(s)).")
                a1, a2 = st.columns(2)
                a1.download_button("⬇️ All-stores summary (.xlsx)", allp["xlsx"],
                                   file_name=f"All Stores - {month} {year} - Summary.xlsx",
                                   mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                   key="dl_all_x", **STRETCH)
                a2.download_button("⬇️ Everything (ZIP)", allp["zip"],
                                   file_name=f"All Stores - {month} {year} - Package.zip",
                                   mime="application/zip", key="dl_all_z", **STRETCH)


# ----------------------------------------------------------------------------
# Build the page
# ----------------------------------------------------------------------------
def build_html_page() -> str:
    rec = get_record(spk)
    files_json = {}
    store_data_json = {}
    if rec:
        cj = rec["checklist_json"]
        files_json[spk] = {cat: [{"id": f.get("id"), "filename": f.get("filename"), "size": f.get("size", 0)}
                                 for f in lst] for cat, lst in (cj.get("files") or {}).items() if lst}
        store_data_json[spk] = rec

    cred_rows = ""
    if role == "admin":
        with _connect() as conn:
            users = conn.execute("SELECT username, name, role, password FROM users ORDER BY role, username").fetchall()
        for u, n, r, p in users:
            args = htmllib.escape(json.dumps([u, n, r, p]), quote=True)
            cred_rows += f"""
        <tr>
            <td class="p-3 font-mono font-medium text-emerald-800">{htmllib.escape(u)}</td>
            <td class="p-3">{htmllib.escape(n)}</td>
            <td class="p-3 uppercase text-[10px] font-bold"><span class="bg-emerald-100 text-emerald-800 px-2 py-0.5 rounded">{htmllib.escape(r)}</span></td>
            <td class="p-3 font-mono">{htmllib.escape(p)}</td>
            <td class="p-3 text-right">
                <button onclick="editUser(...{args})" class="text-emerald-700 font-bold underline mr-2">Edit</button>
                <button onclick="deleteUser({htmllib.escape(json.dumps(u), quote=True)})" class="text-red-600 font-bold underline">Delete</button>
            </td>
        </tr>
    """

    add_user_btn = (
        '<button onclick="openAddUserModal()" class="bg-emerald-700 hover:bg-emerald-800 text-white px-3.5 py-1.5 rounded-xl text-xs font-bold transition shadow-xs">+ Add New User</button>'
        if role == "admin" else ""
    )
    e = lambda s: htmllib.escape(str(s), quote=True)
    return (
        HTML_CODE.replace("REPLACE_CURRENT_NAME", e(st.session_state.name))
        .replace("REPLACE_CURRENT_USERNAME", e(current_token() or ""))
        .replace("REPLACE_CURRENT_ROLE", e(role))
        .replace("REPLACE_ADMIN_HIDDEN_CLASS", "" if role == "admin" else "hidden")
        .replace("REPLACE_CRED_TABLE_ROWS", cred_rows)
        .replace("REPLACE_ACCOUNTANT_HIDDEN_CLASS", "hidden" if role == "accountant" else "")
        .replace("REPLACE_ADD_USER_BTN_HTML", add_user_btn)
        .replace("REPLACE_IS_STORE_MANAGER", "true" if role == "store" else "false")
        .replace("REPLACE_ACTIVE_STORE", e(ctx["store"]))
        .replace("REPLACE_MONTH", e(ctx["month"]))
        .replace("REPLACE_YEAR", e(ctx["year"]))
        .replace("REPLACE_FILES_JSON", json.dumps(files_json).replace("</", "<\\/"))
        .replace("REPLACE_STORE_DATA_JSON", json.dumps(store_data_json, default=str).replace("</", "<\\/"))
    )


# Make sure the sign-in token is in the URL (so a refresh keeps you signed in)
_tok = current_token()
if _tok and query_params.get("auth_user") != _tok:
    st.query_params["auth_user"] = _tok

_flash = st.session_state.pop("flash", None)
if _flash:
    getattr(st, _flash[0], st.info)(_flash[1])
    st.session_state.panel_open = _flash[0] != "success" or st.session_state.get("panel_open", False)

render_panel()
components.html(build_html_page(), height=3000, scrolling=True)
