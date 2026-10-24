"""Report engine for The OZ Store.

Pure pandas/openpyxl (no Streamlit) so it can be tested on its own.
Parses uploaded files and builds the accountant package (Excel + ZIP).
"""
from __future__ import annotations

import io
import re
import zipfile
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Side, Border
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------
# Constants shared with the app
# --------------------------------------------------------------------------
ACCOUNTANT_ORDER = [
    "Beverages", "Dried Flower", "Edibles", "Extracts Ingested", "Extracts Inhaled",
    "Infused Pre-Rolls", "Milled Flower", "Pre-Rolls", "Seeds", "Topicals",
    "Vapes - Disposables", "Vapes - 510 Threads", "Vapes - Pods", "Accessories",
    "Apparel / Swag", "Gift Cards", "Other Income", "Raffles", "System Fees", "Uniforms",
]

CHECKLIST_ITEMS = [
    ("brinks", "Brinks Book"),
    ("green1", "Greenline: Inventory Summary"),
    ("green2", "Greenline: Financials"),
    ("green3", "Greenline: Payment Types"),
    ("mon1", "Moneris: Counter Sales"),
    ("mon2", "Moneris: Online Sales"),
    ("ocs1", "OCS: Account Statement"),
    ("ocs2", "OCS: Invoices"),
    ("gas1", "Gas Cards: Details"),
    ("gas2", "Gas Cards: Summary"),
]

EXPENSE_FIELDS = [
    ("rent", "Store Rent & Lease"), ("salaries", "Employee Salaries"),
    ("utilities", "Utilities"), ("advertising", "Advertising"), ("marketing", "Marketing"),
    ("datadeals", "Data Deals & Analytics"), ("insurance", "Insurance"),
    ("garbage", "Garbage Disposal"), ("security", "Security & POS"),
    ("licenses", "Licenses & Compliance"),
]

# Columns that must never be turned into numbers (IDs, SKUs, names...)
_TEXT_COLS = re.compile(
    r"\b(id|sku|upc|barcode|name|product|description|category|brand|supplier|vendor|store|date|code|plu)\b",
    re.I,
)
MAX_SHEET_ROWS = 50_000

# --------------------------------------------------------------------------
# Small helpers
# --------------------------------------------------------------------------
def parse_money(value) -> float:
    """'$1,234.50', '(200)', '1 234,50' -> float. Blank/invalid -> 0.0."""
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return 0.0 if value != value else float(value)
    s = str(value).strip()
    if not s:
        return 0.0
    neg = (s.startswith("(") and s.endswith(")")) or s.startswith("-")
    s = s.replace(" ", "").replace("\u00a0", "")
    if "," in s and "." not in s and re.search(r",\d{1,2}$", s):
        s = s.replace(",", ".")  # European / Quebec decimal comma
    s = re.sub(r"[^\d.]", "", s)
    try:
        x = float(s)
    except ValueError:
        return 0.0
    return -x if neg else x


def total_expenses(payload: dict) -> float:
    return sum(parse_money(v) for v in (payload.get("expenses") or {}).values())


def missing_docs(payload: dict) -> list[str]:
    chk = payload.get("checklist") or {}
    return [label for key, label in CHECKLIST_ITEMS if not chk.get(key)]


def checklist_done(payload: dict) -> int:
    return len(CHECKLIST_ITEMS) - len(missing_docs(payload))


# --------------------------------------------------------------------------
# Reading uploaded files
# --------------------------------------------------------------------------
def _read_ragged(text: str) -> pd.DataFrame:
    """Fallback for exports with title lines above the header (rows of different widths)."""
    import csv

    rows = list(csv.reader(io.StringIO(text)))
    width = max((len(r) for r in rows), default=0)
    if width == 0:
        return pd.DataFrame()
    rows = [r + [""] * (width - len(r)) for r in rows]
    fill = [sum(1 for v in r if str(v).strip()) for r in rows]
    best = max(fill[:25])
    hdr = next(i for i, n in enumerate(fill[:25]) if n == best)
    df = pd.DataFrame(rows[hdr + 1:], columns=[str(v).strip() or f"Column {j + 1}" for j, v in enumerate(rows[hdr])])
    return df.replace(r"^\s*$", pd.NA, regex=True)


def _read_csv(data: bytes) -> pd.DataFrame:
    last_err = None
    for enc in ("utf-8-sig", "cp1252"):
        try:
            text = data.decode(enc)
        except UnicodeDecodeError as e:
            last_err = e
            continue
        try:
            df = pd.read_csv(io.StringIO(text))
            if df.shape[1] == 1:  # maybe ; or tab separated
                alt = pd.read_csv(io.StringIO(text), sep=None, engine="python")
                if alt.shape[1] > 1:
                    df = alt
            return df
        except (pd.errors.ParserError, pd.errors.EmptyDataError):
            return _read_ragged(text)
    raise ValueError(f"Could not decode CSV: {last_err}")


def read_table(name: str, data: bytes) -> pd.DataFrame | None:
    """CSV / Excel -> DataFrame. Anything else (PDF, images) -> None."""
    ext = name.lower().rsplit(".", 1)[-1] if "." in name else ""
    if ext in ("csv", "txt"):
        df = _read_csv(data)
    elif ext in ("xlsx", "xlsm", "xls"):
        df = pd.read_excel(io.BytesIO(data))
    else:
        return None
    return df.dropna(how="all").dropna(axis=1, how="all").reset_index(drop=True)


def _promote_header(df: pd.DataFrame, needle: str = "category", scan: int = 25) -> pd.DataFrame:
    """Some exports have title lines above the real header; find it."""
    if any(str(c).strip().lower() == needle for c in df.columns):
        return df
    for i in range(min(scan, len(df))):
        row = [str(v).strip() for v in df.iloc[i].tolist()]
        if needle in [r.lower() for r in row]:
            out = df.iloc[i + 1:].copy()
            out.columns = row
            return out.reset_index(drop=True)
    return df


def _coerce_numeric(df: pd.DataFrame) -> pd.DataFrame:
    for c in df.columns:
        if pd.api.types.is_numeric_dtype(df[c]) or _TEXT_COLS.search(str(c)):
            continue
        s = df[c].astype(str).str.strip()
        filled = df[c].notna() & (s != "") & (s.str.lower() != "nan")
        if not filled.any():
            continue
        cleaned = (
            s.str.replace(r"^\((.*)\)$", r"-\1", regex=True)
            .str.replace(r"[$,%\s]", "", regex=True)
        )
        num = pd.to_numeric(cleaned, errors="coerce")
        if num[filled].notna().mean() >= 0.9:
            df[c] = num
    return df


def numeric_cols(df: pd.DataFrame) -> list[str]:
    return [
        c for c in df.columns
        if pd.api.types.is_numeric_dtype(df[c])
        and not pd.api.types.is_bool_dtype(df[c])
        and not _TEXT_COLS.search(str(c))
    ]


def clean_inventory(df: pd.DataFrame) -> tuple[pd.DataFrame, list[str]]:
    """Order categories like the accountant wants, drop subtotal rows, fix numbers."""
    notes: list[str] = []
    df = df.dropna(how="all").dropna(axis=1, how="all").copy()
    df = _promote_header(df)
    df.columns = [str(c).strip() for c in df.columns]

    cat_col = next((c for c in df.columns if c.lower() == "category"), None)
    if cat_col is None:
        notes.append("No 'Category' column found, so rows were kept in their original order.")
        return _coerce_numeric(df), notes
    if cat_col != "Category":
        df = df.rename(columns={cat_col: "Category"})

    cats = df["Category"].astype(str).str.strip()
    drop = cats.str.startswith(">") | cats.str.lower().isin(
        ["", "nan", "none", "total", "totals", "grand total"]
    )
    df, cats = df[~drop].copy(), cats[~drop]

    rank = {c.casefold(): i for i, c in enumerate(ACCOUNTANT_ORDER)}
    order = cats.str.casefold().map(rank)
    unknown = sorted(cats[order.isna()].unique())
    if unknown:
        notes.append("Categories not in the accountant order (placed at the end): " + ", ".join(unknown))
    df["Category"] = cats
    df["_o"] = order.fillna(len(ACCOUNTANT_ORDER))
    df = df.sort_values("_o", kind="stable").drop(columns="_o").reset_index(drop=True)
    return _coerce_numeric(df), notes


# --------------------------------------------------------------------------
# Excel styling helpers
# --------------------------------------------------------------------------
_FONT = "Arial"
_HDR_FILL = PatternFill("solid", fgColor="065F46")
_HDR_FONT = Font(name=_FONT, bold=True, color="FFFFFF", size=10)
_TOT_FILL = PatternFill("solid", fgColor="D1FAE5")
_BAD_FILL = PatternFill("solid", fgColor="FEE2E2")
_BODY = Font(name=_FONT, size=10)
_BOLD = Font(name=_FONT, size=10, bold=True)
_THIN = Border(bottom=Side(style="thin", color="E5E7EB"))
MONEY = '#,##0.00;[Red]-#,##0.00'
INT = '#,##0;[Red]-#,##0'


def _cell_value(v):
    if v is None:
        return None
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    if hasattr(v, "item"):
        return v.item()
    return v


def _title(ws, row: int, text: str, size: int = 12):
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(name=_FONT, bold=True, size=size, color="065F46")


def _header(ws, row: int, headers: list[str]):
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=j, value=h)
        c.font, c.fill = _HDR_FONT, _HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _put(ws, row: int, col: int, value, *, bold=False, fill=None):
    c = ws.cell(row=row, column=col, value=_cell_value(value))
    c.font = _BOLD if bold else _BODY
    c.border = _THIN
    if fill:
        c.fill = fill
    if isinstance(c.value, bool):
        pass
    elif isinstance(c.value, int):
        c.number_format = INT
    elif isinstance(c.value, float):
        c.number_format = MONEY
    return c


def _write_df(ws, df: pd.DataFrame, row: int = 1, totals: bool = True) -> int:
    """Write a styled table; returns the next free row."""
    df = df.head(MAX_SHEET_ROWS)
    _header(ws, row, [str(c) for c in df.columns])
    for i, rec in enumerate(df.itertuples(index=False), start=row + 1):
        for j, v in enumerate(rec, 1):
            _put(ws, i, j, v)
    r = row + len(df) + 1
    nums = numeric_cols(df) if totals else []
    if nums:
        for j, c in enumerate(df.columns, 1):
            if j == 1:
                _put(ws, r, j, "TOTAL", bold=True, fill=_TOT_FILL)
            elif c in nums:
                _put(ws, r, j, float(df[c].sum(skipna=True)), bold=True, fill=_TOT_FILL)
            else:
                _put(ws, r, j, None, fill=_TOT_FILL)
        r += 1
    ws.freeze_panes = ws.cell(row=row + 1, column=1)
    return r + 1


def _block(ws, row: int, title: str, headers: list[str], rows: list[list], bad_when=None) -> int:
    _title(ws, row, title)
    _header(ws, row + 1, headers)
    for i, rec in enumerate(rows, start=row + 2):
        bad = bad_when(rec) if bad_when else False
        for j, v in enumerate(rec, 1):
            _put(ws, i, j, v, fill=_BAD_FILL if bad else None)
    return row + len(rows) + 3


def _autowidth(ws, cap: int = 48):
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is None:
                continue
            n = len(f"{c.value:,.2f}") if isinstance(c.value, float) else len(str(c.value))
            widths[c.column] = max(widths.get(c.column, 0), min(n, cap))
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(10, w + 2)


def _sheet_name(name: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]\*\?/\\:]", "", name).strip()[:31] or "Sheet"
    out, n = base, 2
    while out.lower() in used:
        suffix = f" ({n})"
        out, n = base[: 31 - len(suffix)] + suffix, n + 1
    used.add(out.lower())
    return out


def _to_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _inventory_frame(payload: dict) -> pd.DataFrame | None:
    recs = payload.get("inventory_data")
    return pd.DataFrame(recs) if recs else None


# --------------------------------------------------------------------------
# Report builders
# --------------------------------------------------------------------------
def inventory_workbook(df: pd.DataFrame, title: str = "Inventory") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = _sheet_name(title, set())
    _write_df(ws, df)
    _autowidth(ws)
    return _to_bytes(wb)


def build_store_workbook(
    store: str, period: str, payload: dict, tables: dict[str, tuple[str, pd.DataFrame]]
) -> bytes:
    """One store, one month: Summary + Inventory + one sheet per tabular document."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    used = {"summary"}

    ws["A1"] = f"The OZ Store - {store}"
    ws["A1"].font = Font(name=_FONT, bold=True, size=16, color="065F46")
    ws["A2"] = f"Period: {period}"
    ws["A3"] = f"Generated: {datetime.now():%Y-%m-%d %H:%M}"
    for a in ("A2", "A3"):
        ws[a].font = _BODY
    missing = missing_docs(payload)
    ws["A4"] = ("Missing documents: " + ", ".join(missing)) if missing else "All documents received"
    ws["A4"].font = Font(name=_FONT, bold=True, size=10, color="B91C1C" if missing else "047857")

    files = payload.get("files") or {}
    notes = payload.get("notes") or {}
    chk = payload.get("checklist") or {}
    rows = [
        [label, "Received" if chk.get(key) else "MISSING", (files.get(key) or {}).get("name", ""), notes.get(key, "")]
        for key, label in CHECKLIST_ITEMS
    ]
    r = _block(ws, 6, "Document checklist", ["Document", "Status", "File", "Note"], rows,
               bad_when=lambda rec: rec[1] == "MISSING")

    cooh_text = payload.get("cooh_text", "")
    cooh_amt = parse_money(cooh_text) if str(cooh_text).strip() else None
    r = _block(ws, r, "Cash & chargebacks", ["Item", "Value"], [
        ["Cash on hand", cooh_amt if cooh_amt is not None else "Not entered"],
        ["Cash on hand date", payload.get("cooh_date") or "Not entered"],
        ["Moneris chargebacks", payload.get("chargebacks_text") or "Nothing to report."],
    ])

    exp = payload.get("expenses") or {}
    exp_rows = [[label, parse_money(exp.get(key))] for key, label in EXPENSE_FIELDS]
    exp_rows.append(["TOTAL", total_expenses(payload)])
    r = _block(ws, r, "Operating expenses", ["Expense", "Amount"], exp_rows)

    inv = _inventory_frame(payload)
    if inv is not None:
        nums = numeric_cols(inv)[:12]
        inv_rows = [["Categories", int(inv["Category"].nunique()) if "Category" in inv else len(inv)],
                    ["Line items", len(inv)]]
        inv_rows += [[c if c.lower().startswith("total") else f"Total {c}", float(inv[c].sum(skipna=True))] for c in nums]
        r = _block(ws, r, "Inventory totals", ["Measure", "Value"], inv_rows)
        ws2 = wb.create_sheet(_sheet_name("Inventory", used))
        _write_df(ws2, inv)
        _autowidth(ws2)

    for key, (label, df) in tables.items():
        sheet = wb.create_sheet(_sheet_name(label, used))
        _write_df(sheet, df, totals=True)
        _autowidth(sheet)

    _autowidth(ws)
    return _to_bytes(wb)


def build_consolidated_workbook(period: str, payloads: dict[str, dict]) -> bytes:
    """All stores for one month: Overview, Expenses matrix, All Inventory, Category Totals."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Overview"
    used = {"overview"}
    ws["A1"] = f"The OZ Store - All stores - {period}"
    ws["A1"].font = Font(name=_FONT, bold=True, size=16, color="065F46")
    ws["A2"] = f"Generated: {datetime.now():%Y-%m-%d %H:%M}"
    ws["A2"].font = _BODY

    rows = []
    for store, p in payloads.items():
        cooh = str(p.get("cooh_text", "")).strip()
        rows.append([
            store,
            f"{checklist_done(p)}/{len(CHECKLIST_ITEMS)}",
            parse_money(cooh) if cooh else "Not entered",
            p.get("cooh_date") or "",
            p.get("chargebacks_text") or "",
            total_expenses(p),
            ", ".join(missing_docs(p)) or "-",
        ])
    _block(ws, 4, "Store overview",
           ["Store", "Checklist", "Cash on hand", "Cash date", "Chargebacks", "Total expenses", "Missing"],
           rows, bad_when=lambda rec: rec[6] != "-")
    _autowidth(ws)

    stores = list(payloads)
    wx = wb.create_sheet(_sheet_name("Expenses", used))
    _header(wx, 1, ["Expense"] + stores + ["All stores"])
    for i, (key, label) in enumerate(EXPENSE_FIELDS, start=2):
        vals = [parse_money((payloads[s].get("expenses") or {}).get(key)) for s in stores]
        _put(wx, i, 1, label)
        for j, v in enumerate(vals, 2):
            _put(wx, i, j, v)
        _put(wx, i, len(stores) + 2, sum(vals), bold=True)
    tr = len(EXPENSE_FIELDS) + 2
    _put(wx, tr, 1, "TOTAL", bold=True, fill=_TOT_FILL)
    for j, s in enumerate(stores, 2):
        _put(wx, tr, j, total_expenses(payloads[s]), bold=True, fill=_TOT_FILL)
    _put(wx, tr, len(stores) + 2, sum(total_expenses(payloads[s]) for s in stores), bold=True, fill=_TOT_FILL)
    wx.freeze_panes = "B2"
    _autowidth(wx)

    frames = []
    for s in stores:
        f = _inventory_frame(payloads[s])
        if f is not None:
            f.insert(0, "Store", s)
            frames.append(f)
    if frames:
        allinv = pd.concat(frames, ignore_index=True)
        wa = wb.create_sheet(_sheet_name("All Inventory", used))
        _write_df(wa, allinv, totals=True)
        _autowidth(wa)
        nums = numeric_cols(allinv)
        if "Category" in allinv and nums:
            tot = allinv.groupby("Category", sort=False)[nums].sum()
            rank = {c.casefold(): i for i, c in enumerate(ACCOUNTANT_ORDER)}
            tot = tot.assign(_o=[rank.get(str(c).casefold(), 999) for c in tot.index]) \
                     .sort_values("_o", kind="stable").drop(columns="_o").reset_index()
            wc = wb.create_sheet(_sheet_name("Category Totals", used))
            _write_df(wc, tot, totals=True)
            _autowidth(wc)
    return _to_bytes(wb)


def build_zip(entries: list[tuple[str, bytes]]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for arcname, data in entries:
            z.writestr(arcname, data)
    return buf.getvalue()
